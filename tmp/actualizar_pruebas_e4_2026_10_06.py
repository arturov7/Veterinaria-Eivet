from pathlib import Path
from docx import Document
from docx.shared import Inches

root = Path(__file__).resolve().parents[1]
source = root / "Vedia_Salazar_Monografia_E4_adaptada.docx"
capture = root / "docs/evidencias/rf01_login_invalido_2026-10-06.jpg"
output = root / "tmp/Vedia_Salazar_Monografia_E4_adaptada_2026_10_06.docx"

document = Document(source)
paragraphs = document.paragraphs
assert paragraphs[221].text == "2.8 Pruebas"
assert paragraphs[223].text.startswith("Tabla 13. Casos de prueba ejecutados y pendientes al 03/10/2026")
assert len(document.tables[12].rows) == 8

paragraphs[222].runs[0].text += (
    " El 06/10/2026 se probó el camino de error de RF-01 en el panel web "
    "publicado: un correo y una contraseña ficticios inválidos no permitieron "
    "acceder y el formulario mostró un mensaje de credenciales incorrectas. "
    "También se observó HTTP 401 al llamar crear-propietario sin token. "
    "Estos resultados no acreditan el acceso válido ni RF-02 a RF-04."
)
paragraphs[223].runs[0].text = paragraphs[223].runs[0].text.replace(
    "03/10/2026", "06/10/2026"
)
paragraphs[224].runs[0].text = (
    "La Figura 6 documenta el rechazo de credenciales ficticias en la URL "
    "pública el 06/10/2026. Faltan capturas fechadas de los caminos felices "
    "de RF-01 a RF-04, de sus demás casos de error, de la respuesta de salud "
    "y de los permisos entre cuentas. Se obtendrán con datos ficticios, sin "
    "exponer credenciales ni información clínica. La respuesta online de "
    "salud acredita solo la comprobación de conectividad."
)

table = document.tables[12]
for cells in [
    (
        "CP-08", "Crear propietario sin token", "HTTP 401",
        "HTTP 401; «Inicia sesión» (06/10/2026)",
        "Aprobado solo para la respuesta sin token",
    ),
    (
        "CP-11", "RF-01: login web inválido", "Rechazo y mensaje de error",
        "Sin acceso; mensaje de credenciales incorrectas (06/10/2026)",
        "Aprobado solo para este camino de error",
    ),
]:
    row = table.add_row()
    for cell, value in zip(row.cells, cells):
        cell.text = value

index_anchor = next(p for p in document.paragraphs if p.text == "ÍNDICE DE ANEXOS")
index_entry = index_anchor.insert_paragraph_before(
    "Figura 6. Rechazo de inicio de sesión con credenciales ficticias."
)
index_entry.style = paragraphs[48].style

annex_anchor = next(p for p in document.paragraphs if p.text.startswith("F.3 Ensayo de defensa."))
caption = annex_anchor.insert_paragraph_before(
    "Figura 6. Rechazo de inicio de sesión con credenciales ficticias en "
    "https://eivetveterinaria.netlify.app/, 06/10/2026 a las 16:53 "
    "(UTC−04:00). Fuente: captura propia."
)
caption.style = paragraphs[223].style
image_paragraph = annex_anchor.insert_paragraph_before()
image_paragraph.add_run().add_picture(str(capture), width=Inches(5.8))

document.save(output)
print(output)
