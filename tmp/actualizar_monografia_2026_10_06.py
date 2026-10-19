# -*- coding: utf-8 -*-
"""Actualización localizada de la monografía E4 con evidencia del 06/10/2026."""
from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / "Vedia_Salazar_Monografia_E4_adaptada.docx"
BACKUP = ROOT / "tmp" / "monografia_antes_actualizacion_2026-10-06.docx"
copy2(PATH, BACKUP)
doc = Document(PATH)


def replace_paragraph(index: int, text: str) -> None:
    p = doc.paragraphs[index]
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for run in p.runs[1:]:
        run.text = ""


replacements = {
    57: (
        "La Veterinaria EIVET, en Tarija, requiere organizar información de propietarios, mascotas y atención clínica. "
        "Este trabajo desarrolla un sistema web para el personal veterinario y una aplicación móvil de consulta y solicitud de citas para propietarios. "
        "El proyecto utilizó Kanban para ordenar tareas y Flutter, Supabase y Netlify para la implementación. "
        "El repositorio contiene interfaces, modelo de datos, políticas de seguridad y scripts de comprobación. "
        "El 06/10/2026 la suite local de Flutter terminó con siete pruebas aprobadas. En la web pública se observó el rechazo de credenciales inválidas y la función crear-propietario respondió 401 sin token. "
        "La ruta de salud respondió según lo esperado a GET y HEAD, pero POST devolvió 200 en vez del 405 previsto por el código local. "
        "Treinta solicitudes GET medidas el 03/10/2026 respondieron 200; 29 tardaron tres segundos o menos. Esa medición no evalúa las operaciones clínicas. "
        "La separación entre dos propietarios, varios flujos completos en producción y el instalable final aún requieren comprobación."
    ),
    212: (
        "Integré Supabase Auth para iniciar y cerrar sesión, registrar cuentas móviles y recuperar contraseñas. "
        "El cliente observa los cambios de sesión y asocia las preferencias locales con el identificador del usuario. "
        "El 06/10/2026 el panel web rechazó credenciales ficticias inválidas y permaneció en la pantalla de acceso; el camino feliz y la caducidad o renovación de una sesión en producción siguen sin prueba documentada."
    ),
    216: (
        "Los formularios móviles comprueban campos obligatorios antes de guardar. La función crear-propietario rechaza métodos distintos de POST, JSON inválido, correo con formato incorrecto y contraseñas de menos de ocho caracteres; también exige una sesión válida y rol administrador. "
        "Comprobé el 06/10/2026 que una petición sin token respondió HTTP 401; no se inspeccionaron Auth ni las tablas para confirmar efectos sobre datos. "
        "Las políticas RLS deben impedir el acceso directo no autorizado por PostgREST, incluso si alguien evita la interfaz. Esa defensa sigue sin prueba con dos cuentas. La ruta salud solo devuelve online u offline y no lee datos clínicos."
    ),
    218: (
        "Las contraseñas se entregan a Supabase Auth y no se guardan en las tablas de la aplicación. Las URL públicas del panel y de Supabase emplean HTTPS; comprobé su acceso por esa vía. "
        "Este trabajo no evalúa los mecanismos internos de cifrado del proveedor ni inspecciona todo el tráfico móvil."
    ),
    222: (
        "El 06/10/2026 ejecuté flutter test: siete pruebas locales aprobaron y el proceso terminó con código de salida 0. "
        "Estas pruebas cubren lógica demo y conversión de datos, sin acreditar RLS ni operaciones clínicas de producción. "
        "Repetí scripts/verify_salud.ps1 contra la URL pública: GET y HEAD cumplieron, mientras POST respondió 200 en lugar de 405. "
        "Además, el panel rechazó credenciales ficticias inválidas y crear-propietario respondió 401 sin token. "
        "El reporte local, las salidas de salud y las evidencias de esos dos caminos de error se conservan en docs/evidencias; la matriz completa está en docs/PRUEBAS_Y_EVIDENCIAS.md."
    ),
    223: "Tabla 13. Casos de prueba ejecutados y pendientes al 06/10/2026. Fuente: elaboración propia.",
    224: (
        "Se conserva una captura fechada del rechazo de credenciales inválidas en la web pública. "
        "Siguen pendientes capturas fechadas de los flujos clínicos y de los permisos entre cuentas ficticias, así como la comprobación de efectos sobre Auth y la base de datos. "
        "La respuesta online de salud acredita la consulta de comprobación a PostgreSQL, pero no el funcionamiento de todos los módulos."
    ),
    225: (
        "Limitaciones al 06/10/2026: faltan mediciones del proceso actual de EIVET, acceso externo al tablero privado, verificación remota de RLS con dos propietarios ficticios, caminos felices y de error completos de RF-01 a RF-04, capturas fechadas de los flujos clínicos, prueba integral del CRUD y un APK final instalado. "
        "La medición de la ruta de salud figura en 2.9.1; sus resultados no prueban el rendimiento del CRUD."
    ),
    230: (
        "La comprobación pública del 03/10/2026 confirmó acceso al sitio web y 30 respuestas HTTP 200 de GET /functions/v1/salud. "
        "El 06/10/2026 repetí la prueba de métodos: GET y HEAD cumplieron, pero POST aún devolvió 200 en lugar del 405 esperado por el código local. "
        "No se había sincronizado la Edge Function publicada. Tampoco se confirmó en producción la aplicación de las últimas políticas RLS, el CRUD completo, un APK final ni las pruebas con dos propietarios. El enlace del tablero requiere permiso para terceros."
    ),
    241: (
        "5. Pruebas: la matriz distingue resultados esperados, observados y pendientes (2.8). La suite local de Flutter terminó con siete pruebas aprobadas; la web rechazó credenciales inválidas y crear-propietario devolvió 401 sin token. "
        "En salud, GET y HEAD cumplieron, mientras POST evidenció una discrepancia. Ninguno de esos resultados sustituye las pruebas de los cuatro requisitos Must en producción ni verifica RNF-01."
    ),
    242: (
        "6. Despliegue y documentación: la web pública responde en Netlify y se documentaron pasos de reconstrucción, despliegue y uso (2.9 y anexos). "
        "Un APK final instalado, el despliegue sincronizado de la función de salud y la reproducción del manual por una segunda persona siguen por verificar."
    ),
    278: (
        "Versión web pública: https://eivetveterinaria.netlify.app/. Ruta de salud: https://xbvjhrfgznljtwxghvka.supabase.co/functions/v1/salud. "
        "La apertura pública del panel y 30 respuestas GET de salud se comprobaron el 03/10/2026; el 06/10/2026 se repitió la comprobación de métodos y persistió la discrepancia de POST."
    ),
    291: (
        "F.1 Evidencia disponible. La Figura 1 documenta el tablero Kanban. docs/PRUEBAS_Y_EVIDENCIAS.md contiene la matriz; docs/evidencias/flutter_test_2026-10-06.txt conserva la salida de siete pruebas locales aprobadas. "
        "docs/evidencias/verify_salud_2026-10-06.txt registra GET, HEAD y POST; docs/evidencias/rf01_login_invalido_2026-10-06.jpg muestra el rechazo del login, y docs/evidencias/crear_propietario_sin_token_2026-10-06.txt registra la respuesta 401. "
        "docs/evidencias/rendimiento_salud_2026-10-03.csv conserva las 30 muestras de latencia."
    ),
    292: (
        "F.2 Resultado observado. La suite local aprobó siete pruebas. En producción se observó el rechazo de credenciales inválidas y la respuesta 401 de crear-propietario sin token; no se verificaron efectos sobre datos. "
        "En salud, GET y HEAD cumplieron; POST falló porque devolvió HTTP 200 y el código local espera 405. "
        "La prueba con dos propietarios, los cuatro Must completos, las capturas de los flujos clínicos y la instalación del APK permanecen sin evidencia de aprobación."
    ),
}
for index, value in replacements.items():
    replace_paragraph(index, value)

# Jerarquía coherente para el índice y la navegación de Word.
for p in doc.paragraphs:
    value = p.text.strip()
    if value.startswith(("2.2 ", "2.3 ", "2.4 ", "2.5 ", "2.6 ", "2.7 ", "2.8 ", "2.9 ")):
        p.style = doc.styles.get_by_id("Heading2", 1)
        for run in p.runs:
            run.bold = True
            run.font.size = Pt(11)
    elif value.startswith(("2.3.1 ", "2.3.2 ", "2.3.3 ", "2.3.4 ", "2.4.1 ", "2.4.2 ", "2.4.3 ", "2.4.4 ", "2.7.1 ", "2.7.2 ", "2.7.3 ", "2.7.4 ", "2.7.5 ", "2.9.1 ")):
        p.style = doc.styles.get_by_id("Heading3", 1)
        for run in p.runs:
            run.bold = True
            run.font.size = Pt(11)
    elif value == "Marco de trabajo adoptado: Kanban":
        p.style = doc.styles.get_by_id("Heading3", 1)
    if p.style and p.style.name.startswith("Heading"):
        p.paragraph_format.keep_with_next = True

# La tabla conserva el diseño institucional; se clona la última fila para
# mantener anchos, bordes, rellenos y tipografía.
table = doc.tables[12]


def cell_text(cell, value):
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = value
        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(value)


for cell, value in zip(
    table.rows[7].cells,
    ["CP-07", "Suite Flutter local", "Pruebas sin fallos", "7 aprobadas; salida 0 (06/10/2026)", "Aprobado local"],
):
    cell_text(cell, value)

extra_cases = [
    ["CP-08", "crear-propietario sin token", "HTTP 401", "HTTP 401 (06/10/2026); efectos en BD no inspeccionados", "Aprobado solo para HTTP"],
    ["CP-09", "crear-propietario con veterinario", "HTTP 403; sin alta", "No ejecutado", "Pendiente"],
    ["CP-10", "crear-propietario con administrador", "HTTP 201; alta única", "No ejecutado", "Pendiente"],
    ["CP-11", "RF-01: credenciales inválidas", "Rechazo sin acceso", "Pantalla de acceso y error visible (06/10/2026)", "Aprobado: error"],
    ["CP-12", "RF-02: propietario y mascota", "Alta autorizada; denegación al propietario", "No ejecutado integralmente", "Pendiente"],
    ["CP-13", "RF-03: consulta", "Registro válido; error con mascota inválida", "No ejecutado en producción", "Pendiente"],
    ["CP-14", "RF-04: historial", "Consultas previas; estado vacío", "No ejecutado en producción", "Pendiente"],
]
for values in extra_cases:
    row = deepcopy(table.rows[-1]._tr)
    table._tbl.append(row)
    for cell, value in zip(table.rows[-1].cells, values):
        cell_text(cell, value)

# Repetir la cabecera al continuar la tabla en otra página.
header = table.rows[0]._tr
trPr = header.get_or_add_trPr()
if not trPr.findall(qn("w:tblHeader")):
    trPr.append(OxmlElement("w:tblHeader"))

doc.save(PATH)
print(PATH)
print(BACKUP)
