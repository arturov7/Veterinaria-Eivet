# -*- coding: utf-8 -*-
"""Reconstruye los índices con páginas verificadas en el PDF exportado por Word."""
from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Inches, Pt
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "Vedia_Salazar_Monografia_E4_adaptada.docx"
PDF = Path(sys.argv[1])

reader = PdfReader(PDF)
pages = []
for page in reader.pages:
    lines = [line.strip() for line in (page.extract_text() or "").splitlines() if line.strip()]
    pages.append(lines)


def page_number(prefix: str, exact: bool = False) -> str:
    for lines in pages[9:]:  # después de los preliminares
        if any((line == prefix if exact else line.startswith(prefix)) for line in lines[1:]):
            return lines[0]
    raise ValueError(f"No se encontró la entrada en el cuerpo: {prefix}")


contents = [
    ("Capítulo 1. Introducción", "CAPÍTULO 1. INTRODUCCIÓN"),
    ("1.1 Antecedentes", "1.1 Antecedentes"),
    ("1.2 Descripción del problema", "1.2 Descripción del problema"),
    ("1.3 Planteamiento del problema", "1.3 Planteamiento del problema"),
    ("1.4 Justificación del problema", "1.4 Justificación del problema"),
    ("1.5 Objetivos", "1.5 Objetivos"),
    ("Capítulo 2. Marco teórico", "CAPÍTULO 2. MARCO TEÓRICO"),
    ("2.1 Fundamentos conceptuales y tecnológicos", "2.1 Fundamentos conceptuales y tecnológicos"),
    ("2.2 Metodología de desarrollo", "2.2 Metodología de desarrollo"),
    ("2.3 Análisis de requisitos", "2.3 Análisis de requisitos"),
    ("2.4 Diseño de la solución", "2.4 Diseño de la solución"),
    ("2.5 Justificación del stack tecnológico", "2.5 Justificación del stack tecnológico"),
    ("2.6 Implementación", "2.6 Implementación"),
    ("2.7 Seguridad", "2.7 Seguridad"),
    ("2.8 Pruebas", "2.8 Pruebas"),
    ("2.9 Despliegue", "2.9 Despliegue"),
    ("Capítulo 3. Conclusiones y recomendaciones", "CAPÍTULO 3. CONCLUSIONES Y RECOMENDACIONES"),
    ("3.1 Conclusiones", "3.1 Conclusiones"),
    ("3.2 Recomendaciones", "3.2 Recomendaciones"),
    ("Bibliografía consultada", "BIBLIOGRAFÍA CONSULTADA"),
    ("Anexos", "ANEXOS"),
]
tables = [
    "Soluciones informáticas existentes en el dominio",
    "Puntos críticos y datos publicados de otras veterinarias",
    "Correspondencia objetivo, evidencia y conclusión",
    "Periodos de revisión Kanban y evidencias del repositorio",
    "Actores del sistema",
    "Requisitos funcionales, criterios de aceptación y prioridad",
    "Requisitos no funcionales",
    "Capacidades declaradas fuera de alcance",
    "Componentes y responsabilidades",
    "Diccionario de datos de las entidades principales",
    "Operaciones utilizadas por Flutter y Edge Functions",
    "Stack tecnológico y despliegue real",
    "Casos de prueba ejecutados y pendientes",
]
figures = [
    "Tablero Kanban de EIVET",
    "Casos de uso y permisos por actor",
    "Componentes y protocolos de EIVET",
    "Entidades principales y referencias del esquema",
    "Wireframes iniciales de la aplicación móvil",
]
annexes = [
    ("Anexo A. Manual de usuario", "Anexo A. MANUAL DE USUARIO"),
    ("Anexo B. Manual de instalación y despliegue", "Anexo B. MANUAL DE INSTALACIÓN Y DESPLIEGUE"),
    ("Anexo C. Repositorio y tablero", "Anexo C. REPOSITORIO Y TABLERO"),
    ("Anexo D. URL y paquete instalable", "Anexo D. URL Y PAQUETE INSTALABLE"),
    ("Anexo E. Diagramas y trazabilidad", "Anexo E. DIAGRAMAS Y TRAZABILIDAD"),
    ("Anexo F. Evidencias, casos de prueba y defensa", "Anexo F. EVIDENCIAS, CASOS DE PRUEBA Y DEFENSA"),
]

doc = Document(DOCX)


def heading(text):
    return next(p for p in doc.paragraphs if p.text.strip() == text)


def remove_between(start, end):
    paras = doc.paragraphs
    a = next(i for i, p in enumerate(paras) if p._p is start._p)
    b = next(i for i, p in enumerate(paras) if p._p is end._p)
    for p in paras[a + 1 : b]:
        p._element.getparent().remove(p._element)


def format_entry(p, label, number):
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.tab_stops.add_tab_stop(Inches(5.68), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p.add_run(f"{label}\t{number}")
    for run in p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(11)


content_h = heading("ÍNDICE DE CONTENIDO")
table_h = heading("ÍNDICE DE TABLAS")
figure_h = heading("ÍNDICE DE FIGURAS")
annex_h = heading("ÍNDICE DE ANEXOS")
summary_h = heading("RESUMEN")

remove_between(content_h, table_h)
remove_between(table_h, figure_h)
remove_between(figure_h, annex_h)
remove_between(annex_h, summary_h)

for label, anchor in contents:
    format_entry(table_h.insert_paragraph_before(), label, page_number(anchor, exact=True))
for n, label in enumerate(tables, 1):
    format_entry(figure_h.insert_paragraph_before(), f"Tabla {n}. {label}", page_number(f"Tabla {n}."))
for n, label in enumerate(figures, 1):
    format_entry(annex_h.insert_paragraph_before(), f"Figura {n}. {label}", page_number(f"Figura {n}."))
for label, anchor in annexes:
    format_entry(summary_h.insert_paragraph_before(), label, page_number(anchor, exact=True))

doc.save(DOCX)
print("Índices actualizados:", len(contents), len(tables), len(figures), len(annexes))
