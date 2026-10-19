"""Actualiza el perfil EIVET para E3 conservando el archivo de origen."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

SOURCE = Path('Vedia_Salazar_Avance_Capitulo2_Implementacion_Seguridad_Pruebas.docx')
OLDER = Path('Vedia_Salazar_Perfil_Proyecto_Corregido.docx')
OUTPUT = Path('Vedia_Salazar_Perfil_Proyecto_E3_corregido.docx')
TMP = Path('tmp/perfil_e3')
TMP.mkdir(parents=True, exist_ok=True)
doc = Document(SOURCE)
orig = list(doc.paragraphs)


def font(size=26, bold=False):
    base = Path('C:/Windows/Fonts')
    name = 'arialbd.ttf' if bold else 'arial.ttf'
    return ImageFont.truetype(str(base / name), size)


def box(draw, rect, title, lines):
    draw.rounded_rectangle(rect, radius=16, fill='#f1f7f3', outline='#15513e', width=3)
    x1, y1, x2, _ = rect
    draw.text((x1 + 17, y1 + 12), title, font=font(29, True), fill='#083e2d')
    y = y1 + 53
    for line in lines:
        draw.text((x1 + 17, y), line, font=font(22), fill='#172b25')
        y += 33


def arrow(draw, a, b, label='', label_xy=None):
    draw.line((a, b), fill='#19785c', width=5)
    from math import atan2, cos, sin, pi
    angle = atan2(b[1]-a[1], b[0]-a[0])
    for delta in (-pi/6, pi/6):
        tip = (b[0]-20*cos(angle+delta), b[1]-20*sin(angle+delta))
        draw.line((b, tip), fill='#19785c', width=5)
    if label:
        x, y = label_xy or ((a[0]+b[0])/2, (a[1]+b[1])/2)
        draw.text((int(x), int(y)), label, font=font(19, True), fill='#144b3b')


arch = Image.new('RGB', (1600, 1060), 'white')
d = ImageDraw.Draw(arch)
d.text((430, 18), 'Arquitectura tecnológica de EIVET', font=font(39, True), fill='#063e2c')
box(d, (50, 105, 475, 260), 'Panel del personal', ['Flutter Web · Netlify', 'Navegador por HTTPS'])
box(d, (50, 350, 475, 505), 'App del propietario', ['Flutter Android · APK', 'Sesión Supabase Auth'])
box(d, (680, 95, 1110, 225), 'Supabase Auth', ['Identidad y JWT'])
box(d, (680, 300, 1110, 430), 'PostgREST', ['API de tablas · RLS'])
box(d, (680, 505, 1110, 650), 'Edge Functions', ['crear-propietario · salud'])
box(d, (680, 725, 1110, 870), 'Supabase Storage', ['Bucket avatars · RLS'])
box(d, (1260, 290, 1560, 540), 'PostgreSQL', ['Tablas públicas', 'Políticas RLS', 'RPC salud()'])
arrow(d, (475, 180), (680, 160), 'HTTPS / JSON', (505, 125))
arrow(d, (475, 420), (680, 365), 'HTTPS / JSON + JWT', (492, 337))
arrow(d, (475, 220), (680, 555), 'HTTPS / JSON + JWT', (465, 245))
arrow(d, (475, 465), (680, 770), 'HTTPS / API Storage', (440, 610))
arrow(d, (1110, 365), (1260, 365), 'SQL interno', (1102, 315))
arrow(d, (1110, 570), (1260, 470), 'RPC / SQL', (1120, 510))
d.text((55, 950), 'La clave de servicio solo se usa dentro de la función crear-propietario.', font=font(23), fill='#333333')
arch_path = TMP / 'arquitectura_e3.png'
arch.save(arch_path)

erd = Image.new('RGB', (1700, 1420), 'white')
d = ImageDraw.Draw(erd)
d.text((440, 15), 'Modelo de datos principal de EIVET', font=font(39, True), fill='#063e2c')
box(d, (650, 85, 1060, 185), 'auth.users', ['id UUID · Auth'])
box(d, (35, 260, 485, 420), 'clientes', ['id PK/FK auth.users', 'nombre, correo'])
box(d, (630, 260, 1080, 420), 'mascotas', ['id PK', 'cliente_id o propietario_id FK'])
box(d, (1215, 260, 1665, 420), 'propietarios', ['id PK · ficha manual', 'nombre_completo, correo'])
box(d, (35, 550, 485, 710), 'citas', ['id PK · mascota_id FK', 'cliente_id / propietario_id'])
box(d, (630, 550, 1080, 710), 'consultas', ['id PK · mascota_id FK', 'veterinario_id FK perfiles'])
box(d, (1215, 550, 1665, 710), 'vacunas', ['id PK · mascota_id FK', 'fecha_aplicacion, proxima_dosis'])
box(d, (35, 850, 485, 1010), 'historiales_clinicos', ['id PK · mascota_id FK', 'autorizado, diagnostico'])
box(d, (630, 850, 1080, 1010), 'tratamientos', ['id PK · mascota_id FK', 'consulta_id FK opcional'])
box(d, (1215, 850, 1665, 1010), 'perfiles', ['id PK/FK auth.users', 'rol: admin / veterinario'])
arrow(d, (850, 185), (260, 260), '1:1', (490, 185))
arrow(d, (850, 185), (1440, 850), '1:1', (1450, 805))
arrow(d, (485, 340), (630, 340), '1:N', (520, 300))
arrow(d, (1215, 340), (1080, 340), '1:N', (1115, 300))
arrow(d, (720, 420), (350, 550), '1:N', (470, 455))
arrow(d, (850, 420), (850, 550), '1:N', (865, 455))
arrow(d, (1000, 420), (1350, 550), '1:N', (1115, 460))
arrow(d, (660, 420), (330, 850), '1:N', (390, 760))
arrow(d, (925, 710), (925, 850), '0:N', (940, 765))
d.text((55, 1120), 'Cada mascota tiene un solo titular: cliente_id o propietario_id (script 12).', font=font(26, True), fill='#103e31')
d.text((55, 1170), 'Tratamientos también referencia mascota_id. Vacunas no referencia consultas.', font=font(25), fill='#333333')
d.text((55, 1220), 'Otras tablas del esquema: servicios y notificaciones. auth.users pertenece a Supabase Auth.', font=font(25), fill='#333333')
erd_path = TMP / 'modelo_e3.png'
erd.save(erd_path)

uses = Image.new('RGB', (1700, 1090), 'white')
d = ImageDraw.Draw(uses)
d.text((430, 18), 'Casos de uso y permisos de EIVET', font=font(38, True), fill='#063e2c')
box(d, (40, 125, 485, 285), 'Administrador', ['Autenticado en el panel', 'Rol en public.perfiles'])
box(d, (40, 400, 485, 560), 'Veterinario', ['Autenticado en el panel', 'Rol en public.perfiles'])
box(d, (40, 680, 485, 840), 'Propietario', ['Cuenta en public.clientes', 'Acceso móvil'])
box(d, (690, 95, 1640, 235), 'CU-01 Iniciar sesión', ['Administrador, veterinario y propietario'])
box(d, (690, 275, 1640, 415), 'CU-02 Registrar o editar mascota', ['Solo administrador y veterinario'])
box(d, (690, 455, 1640, 595), 'CU-03 Registrar consulta', ['Solo administrador y veterinario'])
box(d, (690, 635, 1640, 775), 'RF-02 Crear cuenta de propietario', ['Solo administrador; Edge Function'])
box(d, (690, 815, 1640, 955), 'RF-06 y RF-08 Consulta y citas', ['Propietario: datos propios y citas propias'])
arrow(d, (485, 200), (690, 165))
arrow(d, (485, 220), (690, 340))
arrow(d, (485, 235), (690, 525))
arrow(d, (485, 255), (690, 705))
arrow(d, (485, 465), (690, 365))
arrow(d, (485, 485), (690, 550))
arrow(d, (485, 740), (690, 910))
d.text((50, 1020), 'El propietario no registra mascotas ni modifica datos clínicos.', font=font(25, True), fill='#154c3b')
uses_path = TMP / 'casos_uso_e3.png'
uses.save(uses_path)

# Versiones legibles para insertar a ancho de página: sin flechas sobre texto.
arch = Image.new('RGB', (1600, 1080), 'white')
d = ImageDraw.Draw(arch)
d.text((415, 22), 'Arquitectura tecnológica de EIVET', font=font(39, True), fill='#063e2c')
box(d, (50, 115, 750, 280), 'Panel del personal', ['Flutter Web servido por Netlify', 'Navegador: HTTPS / HTML + JS'])
box(d, (850, 115, 1550, 280), 'App del propietario', ['Flutter Android · APK', 'Sesión de usuario de Supabase'])
arrow(d, (400, 280), (650, 455), 'HTTPS / JSON + JWT', (260, 330))
arrow(d, (1200, 280), (950, 455), 'HTTPS / JSON + JWT', (1050, 330))
box(d, (340, 455, 1260, 755), 'Servicios administrados de Supabase', [
    'Auth: identidad, sesiones y JWT',
    'PostgREST: API CRUD con políticas RLS',
    'Edge Functions: crear-propietario y salud',
    'Storage: bucket avatars y políticas de acceso',
])
arrow(d, (800, 755), (800, 840), 'SQL interno / RPC', (830, 775))
box(d, (470, 840, 1130, 1000), 'PostgreSQL', ['Tablas, claves y restricciones', 'RLS y función salud()'])
arch.save(arch_path)

erd = Image.new('RGB', (1700, 1530), 'white')
d = ImageDraw.Draw(erd)
d.text((420, 20), 'Modelo de datos principal de EIVET', font=font(39, True), fill='#063e2c')
box(d, (600, 90, 1100, 210), 'auth.users', ['id UUID · Supabase Auth'])
box(d, (40, 270, 510, 430), 'clientes', ['id PK/FK auth.users', 'nombre, correo'])
box(d, (615, 270, 1085, 430), 'mascotas', ['id PK', 'cliente_id / propietario_id FK'])
box(d, (1190, 270, 1660, 430), 'propietarios', ['id PK · ficha manual', 'nombre_completo, correo'])
box(d, (40, 515, 510, 675), 'citas', ['id PK · mascota_id FK', 'cliente_id / propietario_id'])
box(d, (615, 515, 1085, 675), 'consultas', ['id PK · mascota_id FK', 'veterinario_id FK perfiles'])
box(d, (1190, 515, 1660, 675), 'vacunas', ['id PK · mascota_id FK', 'fecha_aplicacion, proxima_dosis'])
box(d, (40, 760, 510, 920), 'historiales_clinicos', ['id PK · mascota_id FK', 'autorizado, diagnostico'])
box(d, (615, 760, 1085, 920), 'tratamientos', ['id PK · mascota_id FK', 'consulta_id FK opcional'])
box(d, (1190, 760, 1660, 920), 'perfiles', ['id PK/FK auth.users', 'rol: administrador / veterinario'])
d.text((45, 1020), 'RELACIONES DEL ESQUEMA', font=font(30, True), fill='#063e2c')
relations = [
    'auth.users.id → clientes.id y perfiles.id',
    'mascotas.cliente_id → clientes.id; mascotas.propietario_id → propietarios.id',
    'citas, consultas, vacunas, historiales y tratamientos → mascotas.id',
    'tratamientos.consulta_id → consultas.id (opcional)',
    'consultas.veterinario_id → perfiles.id (opcional)',
    'Cada mascota tiene un solo titular según el script 12.',
    'Tablas complementarias: servicios y notificaciones.',
]
for i, line in enumerate(relations):
    d.text((45, 1075 + i*52), line, font=font(24), fill='#263831')
erd.save(erd_path)


def setp(index, text, style=None):
    p = orig[index]
    p.text = text
    if style:
        p.style = style
    return p


def add_after(anchor, text='', style=None):
    p = doc.add_paragraph(text, style=style)
    anchor._element.addnext(p._element)
    return p


def move_table_after(table, anchor):
    anchor._element.addnext(table._element)


def fill_table(table, rows, size=8):
    while table.rows:
        table._tbl.remove(table.rows[0]._tr)
    for values in rows:
        cells = table.add_row().cells
        for c, value in zip(cells, values):
            c.text = value
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(size)
    table.autofit = True


def replace_picture(index, path, width=15.4):
    p = orig[index]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width))


# Portada, delimitación y antecedentes.
setp(4, 'Módulo 4 · Entregable E3 · Checkpoint técnico · 3 de octubre de 2026')
setp(32, 'Desarrollar una aplicación web y móvil multiplataforma para gestionar propietarios, mascotas, consultas y citas de la Veterinaria EIVET de la ciudad de Tarija, mediante un panel para el personal autorizado y una aplicación de consulta y solicitud de citas para los propietarios.')
setp(14, '1.1 Antecedentes', 'Heading 2')
p = orig[14]
p = add_after(p, 'La gestión veterinaria reúne datos de pacientes animales, titulares, consultas, vacunas y agenda. Para revisar soluciones existentes se consultaron las páginas oficiales de VetCloud, DaySmart Vet y VetiCloud el 3 de octubre de 2026. Estas describen historias clínicas y agenda; algunas incluyen facturación, inventario o colaboración entre clínicas. La comparación se limita a las funciones declaradas por cada proveedor, sin afirmar resultados de implantación en EIVET.')
p = add_after(p, 'EIVET necesita un alcance local más acotado: panel web para administrador y veterinario, consulta móvil de la mascota por su propietario y solicitudes de citas, con datos compartidos en Supabase. La selección tecnológica y las restricciones se detallan en los requisitos y en el Capítulo 2. Las mediciones del proceso manual todavía requieren observación directa en la clínica.')

fill_table(doc.tables[1], [
    ('Solución', 'Alcance documentado', 'Qué resuelve', 'Qué no cubre para EIVET', 'Fuente consultada'),
    ('VetCloud', 'Gestión clínica, fichas y agenda.', 'Centraliza pacientes y recordatorios.', 'No acredita una adaptación al proceso específico de EIVET; habría que evaluarla.', 'VetCloud, https://vetcloud.cl/ (03/10/2026).'),
    ('DaySmart Vet', 'Historia clínica, citas, comunicación, pagos e inventario.', 'Integra agenda y registros en la nube.', 'Incluye procesos comerciales fuera del alcance EIVET; no se evaluó su implantación local.', 'DaySmart Vet, https://www.daysmart.com/vet/ (03/10/2026).'),
    ('VetiCloud', 'Historia clínica compartida, visitas y app para dueños.', 'Conecta clínicas y propietarios.', 'Su enfoque multicentro supera el alcance de una clínica; no se evaluó su implantación local.', 'VetiCloud, https://veticloud.com/es (03/10/2026).'),
], 7)
p._element.addnext(doc.tables[1]._element)
doc.tables[1]._element.addnext(orig[11]._element)
for blank in (orig[12], orig[13]):
    blank._element.getparent().remove(blank._element)

# Tabla 3: ninguna cifra inventada.
fill_table(doc.tables[2], [
    ('N.º', 'Punto crítico del proceso actual', 'Dato del contexto / medición'),
    ('1', 'Tiempo para localizar una ficha de mascota', 'Pendiente de medición en EIVET: cronometrar al menos 10 búsquedas reales y registrar mediana y rango.'),
    ('2', 'Tiempo de registro de una consulta', 'Pendiente de medición en EIVET: observar al menos 10 consultas y registrar inicio, fin y medio usado.'),
    ('3', 'Registros duplicados o incompletos', 'Pendiente de medición en EIVET: revisar muestra autorizada de fichas sin copiar datos personales.'),
], 8)

# Capítulo 2 y fundamentos, antes de la metodología.
method = orig[46]
intro = doc.add_paragraph('CAPÍTULO 2. MARCO TEÓRICO Y DESARROLLO', style='Heading 1')
method._element.addprevious(intro._element)
h = doc.add_paragraph('2.1 Fundamentos conceptuales y tecnológicos', style='Heading 2')
method._element.addprevious(h._element)
for text in [
    'Un expediente veterinario digital relaciona la identidad del propietario con la mascota y sus atenciones, vacunas y tratamientos. La integridad referencial evita que una cita o consulta apunte a una mascota inexistente. En EIVET, el historial visible al propietario depende de la autorización definida en el esquema y de las políticas de acceso, no solo de filtros visuales.',
    'Flutter 3.44.8 y Dart 3.12.2 permiten mantener una base de código para la interfaz Android y el panel web; este proyecto selecciona la vista mediante kIsWeb. La aplicación web se publica en Netlify y el APK se genera por separado. La documentación de Flutter describe el soporte web desde la misma base de código (Flutter, s. f.).',
    'Supabase Auth gestiona identidades y sesiones mediante JWT; PostgREST expone operaciones CRUD sobre PostgreSQL por HTTPS/JSON. Las políticas Row Level Security (RLS) se evalúan en la base de datos para limitar filas por cuenta y rol. Supabase Storage aloja avatares, y las Edge Functions encapsulan operaciones que requieren lógica de servidor, como crear cuentas de propietario (Supabase, s. f.).',
    'Kanban representa tareas en un flujo visible y limita el trabajo en curso. En este proyecto las columnas Hacer, En curso y Hecho permiten priorizar el desarrollo individual; el límite acordado es dos tarjetas simultáneas. Las revisiones por fecha sirven para documentar avances, sin convertir el flujo en Scrum ni afirmar sprints no realizados (Kanban Guides, 2025).',
]:
    p = doc.add_paragraph(text)
    method._element.addprevious(p._element)
setp(46, '2.2 Metodología de desarrollo', 'Heading 1')
setp(47, 'Marco de trabajo adoptado: Kanban', 'Heading 2')
setp(49, 'El tablero se revisó al cierre de periodos de trabajo para mover tarjetas según evidencia del repositorio. El flujo fue continuo; no se realizaron ceremonias ni sprints de Scrum. Se limita En curso a dos tarjetas, como muestra la captura conservada de GitHub Projects.')
setp(52, 'Herramienta: GitHub Projects · Tablero: https://github.com/users/arturov7/projects/2/views/1?layout=board · Límite de trabajo en curso: 2 tarjetas. El enlace devuelve 404 sin sesión; debe verificarse el acceso del docente con la cuenta autorizada.')
setp(53, '')
older = Document(OLDER)
old_p = older.paragraphs[49]
rid = old_p._p.xpath('.//a:blip')[0].get(qn('r:embed'))
board_path = TMP / 'tablero_kanban.png'
board_path.write_bytes(older.part.rels[rid].target_part.blob)
replace_picture(53, board_path, 15.5)
setp(54, 'La captura procede de la versión anterior del perfil. La imagen no muestra fecha en pantalla; el estado representado debe cotejarse con GitHub Projects.')
setp(55, 'Figura 1. Tablero Kanban de EIVET con tareas en Hacer, En curso y Hecho. Fuente: captura del proyecto conservada de la versión anterior del perfil.')
fill_table(doc.tables[4], [
    ('Iteración de revisión', 'Periodo documentado', 'Trabajo observado', 'Evidencia / estado'),
    ('1', '31 ago – 9 sep 2026', 'Base móvil y web del proyecto.', 'Commits a807945 y aababb7; realizado.'),
    ('2', '22 – 24 sep 2026', 'Estructura y panel web; vinculación de clientes, mascotas y citas.', 'Commits 8e27770, ed2c194 y 01f9cbd; realizado.'),
    ('3', '25 – 26 sep 2026', 'Creación de cuentas, correcciones móviles y configuración de Netlify.', 'Commits b7e8f63 a d8137d4; publicación web observada.'),
    ('4', '3 oct 2026', 'Ajuste de permisos, ruta de salud y pruebas documentadas.', 'Commit a1a0acb; pruebas RLS de dos cuentas pendientes.'),
], 8)

# Actores y requisitos: distinguir administrador, veterinario y propietario.
fill_table(doc.tables[5], [
    ('ID', 'Actor', 'Descripción', 'Puede realizar', 'Restricción'),
    ('A-01', 'Administrador', 'Personal autenticado con rol administrador en perfiles.', 'Crear cuentas de propietario; gestionar datos clínicos y administrativos.', 'No usa la clave de servicio desde el navegador.'),
    ('A-02', 'Veterinario', 'Personal autenticado con rol veterinario.', 'Registrar y actualizar mascotas, consultas, tratamientos, vacunas y citas.', 'La función crear-propietario exige rol administrador.'),
    ('A-03', 'Propietario', 'Cuenta de Auth vinculada a clientes.id.', 'Consultar sus mascotas y datos autorizados; solicitar y cancelar sus citas.', 'No registra, edita ni elimina mascotas; no modifica datos clínicos.'),
], 7)
rf2 = doc.tables[6].rows[2]
rf2.cells[1].text = 'Como personal autorizado, quiero registrar una mascota vinculada a un solo titular; el administrador puede crear la cuenta de acceso del propietario mediante la función segura.'
rf2.cells[2].text = 'CA-02.1: administrador crea la cuenta con datos válidos; personal registra la mascota con cliente_id o propietario_id. CA-02.2: se rechaza una mascota sin titular o con dos titulares; el propietario no puede escribirla.'
doc.tables[6].rows[1].cells[1].text = 'Como usuario registrado, quiero iniciar sesión para acceder solo a las funciones autorizadas para mi perfil.'
doc.tables[3].rows[6].cells[2].text = '2.6 Implementación y README de despliegue; APK aún sin evidencia de distribución.'
setp(87, 'CU-02 Registrar mascota', 'Heading 2')
orig[86]._element.getparent().remove(orig[86]._element)
setp(88, 'Actor principal: administrador o veterinario autenticado. Solo el administrador puede crear una cuenta de acceso nueva.')
setp(89, 'Precondición: existe una cuenta en clientes o una ficha manual en propietarios; si se necesita cuenta nueva, el administrador la crea primero.')
setp(90, 'El personal abre Mascotas en el panel web y elige registrar.')
setp(91, 'El sistema muestra el formulario de mascota y solicita seleccionar exactamente un titular.')
setp(92, 'El personal completa nombre y demás datos disponibles.')
setp(93, 'El sistema valida los campos obligatorios y el vínculo con el titular.')
setp(94, 'El personal confirma el registro de la mascota.')
setp(95, 'PostgREST almacena la mascota si el rol y las políticas RLS lo permiten.')
setp(96, 'Postcondición: la mascota queda vinculada al titular seleccionado.')
setp(98, 'A1. Si la cuenta no existe, el administrador la crea mediante crear-propietario o selecciona una ficha manual autorizada.')
setp(99, 'A2. Si se seleccionan dos titulares o ninguno, el registro se rechaza. Si el usuario es propietario, la escritura debe ser denegada por RLS.')
setp(100, '')
replace_picture(71, uses_path)
setp(72, 'Figura 2. Casos de uso y permisos por actor. Fuente: elaboración propia según RF-01 a RF-08 y función crear-propietario.')

# Arquitectura y modelo según código y SQL, sin presentar migraciones como aplicadas.
setp(117, 'Arquitectura cliente–servicio administrado con Flutter, Supabase y Netlify.', 'Heading 3')
setp(118, 'El navegador recibe Flutter Web desde Netlify por HTTPS. El teléfono ejecuta Flutter Android. Ambos clientes consumen los servicios de Supabase por HTTPS/JSON; no existe un servidor de aplicación propio.')
setp(119, 'Supabase Auth emite sesiones; PostgREST expone tablas; Edge Functions ejecuta crear-propietario y salud; PostgreSQL aplica relaciones, restricciones y las políticas RLS instaladas.', 'Normal')
setp(121, 'PostgREST conecta la aplicación con PostgreSQL. El repositorio contiene políticas RLS para autorización por fila; su estado remoto debe verificarse con pruebas de dos cuentas antes de dar por satisfecho RNF-02.')
setp(122, 'PostgreSQL almacena clientes, perfiles, propietarios, mascotas, consultas, tratamientos, vacunas, historiales y citas mediante relaciones. Las reglas de los scripts deben aplicarse en el proyecto remoto para que limiten los registros visibles.')
setp(124, 'La figura etiqueta los intercambios HTTPS/JSON, JWT y SQL interno. La clave de servicio permanece en la Edge Function; la aplicación usa la publishable key.')
replace_picture(125, arch_path)
setp(126, 'Figura 3. Componentes y protocolos de EIVET. Fuente: elaboración propia a partir de lib/main.dart, supabase/config.toml y netlify.toml.')
fill_table(doc.tables[9], [
    ('Componente', 'Responsabilidad', 'Tecnología', 'Se comunica con', 'Protocolo / formato'),
    ('Panel del personal', 'Interfaz administrativa.', 'Flutter Web 3.44.8; Netlify', 'Navegador; Supabase', 'HTTPS/HTML/JS; HTTPS/JSON + JWT'),
    ('App del propietario', 'Consulta y citas.', 'Flutter Android 3.44.8', 'Supabase', 'HTTPS/JSON + JWT'),
    ('Identidad', 'Sesiones y JWT.', 'Supabase Auth', 'Clientes; auth.users', 'HTTPS/JSON; JWT'),
    ('API de datos', 'CRUD de tablas expuestas.', 'Supabase PostgREST', 'Flutter; PostgreSQL', 'HTTPS/JSON; SQL interno'),
    ('Funciones', 'Alta segura de cuentas y salud.', 'Supabase Edge Functions', 'Auth; RPC salud', 'HTTPS/JSON; SQL interno'),
    ('Persistencia', 'Datos y restricciones por fila.', 'PostgreSQL + RLS', 'PostgREST; funciones', 'SQL interno'),
    ('Archivos', 'Avatares y políticas del bucket.', 'Supabase Storage', 'App Flutter', 'HTTPS/API Storage'),
], 7)
setp(128, '2.4.2 Modelo de datos', 'Heading 2')
model_note = add_after(orig[128], 'El esquema combina cuentas móviles en clientes (id = auth.users.id) y fichas manuales en propietarios. Cada mascota referencia un solo titular mediante cliente_id o propietario_id. Citas conserva el vínculo con mascota y titular; consultas, historiales_clinicos, vacunas y tratamientos referencian mascota. Las vacunas no tienen consulta_id; tratamientos sí admite una consulta_id opcional. Perfiles almacena los roles del personal; nunca guarda contraseñas. El diagrama muestra las entidades centrales y deja servicios y notificaciones como tablas complementarias.')
model_note = add_after(model_note, 'Las políticas previstas en los scripts 12 y 13 limitan escritura de mascotas al personal y lectura del propietario a cliente_id = auth.uid(), además de validar la cita contra la mascota. Los scripts están en el repositorio; la ejecución remota de estas últimas migraciones y la prueba de dos cuentas aún no están documentadas como aprobadas.')
replace_picture(129, erd_path)
setp(130, 'Figura 4. Entidades principales y referencias del esquema de Supabase. Fuente: elaboración propia a partir de los scripts SQL del proyecto.')
setp(131, 'Tabla 11. Diccionario de datos de las entidades principales. Fuente: scripts SQL del proyecto.')
setp(132, '')
fill_table(doc.tables[10], [
    ('Entidad', 'Campo', 'Tipo', 'Obligatorio', 'Clave / referencia', 'Descripción'),
    ('clientes', 'id', 'UUID', 'Sí', 'PK; FK auth.users.id', 'Cuenta móvil autenticada.'),
    ('clientes', 'nombre, correo', 'TEXT', 'Nombre sí; correo según script 10', '—', 'Datos del propietario con acceso.'),
    ('perfiles', 'id, rol', 'UUID, TEXT', 'Sí', 'PK/FK auth.users; CHECK rol', 'Personal administrador o veterinario.'),
    ('propietarios', 'id, nombre_completo', 'UUID, TEXT', 'Sí', 'PK', 'Ficha manual, sin cuenta de Auth.'),
    ('mascotas', 'id, nombre', 'UUID, TEXT', 'Sí', 'PK', 'Paciente animal.'),
    ('mascotas', 'cliente_id / propietario_id', 'UUID', 'Uno de dos', 'FK clientes / propietarios', 'Titular único con script 12.'),
    ('mascotas', 'especie, raza, sexo, fecha_nacimiento, peso', 'TEXT, DATE, NUMERIC', 'No', '—', 'Datos descriptivos.'),
    ('citas', 'id, mascota_id', 'UUID', 'Sí', 'PK; FK mascotas', 'Solicitud o atención agendada.'),
    ('citas', 'cliente_id / propietario_id', 'UUID', 'Según titular', 'FK clientes / propietarios', 'Titular relacionado.'),
    ('citas', 'fecha_hora, motivo, estado', 'TIMESTAMPTZ, TEXT', 'Sí', 'CHECK estado', 'Agenda y estado.'),
    ('consultas', 'id, mascota_id, veterinario_id', 'UUID', 'Mascota sí', 'PK; FK mascotas / perfiles', 'Atención del panel.'),
    ('consultas', 'fecha, motivo, diagnostico', 'DATE, TEXT', 'Fecha y motivo sí', '—', 'Datos de consulta.'),
    ('historiales_clinicos', 'id, mascota_id, autorizado', 'UUID, BOOLEAN', 'Sí', 'PK; FK mascotas', 'Historial visible según autorización.'),
    ('vacunas', 'id, mascota_id, nombre', 'UUID, TEXT', 'Sí', 'PK; FK mascotas', 'Vacuna registrada.'),
    ('vacunas', 'fecha_aplicacion, proxima_dosis', 'DATE', 'No', '—', 'Fechas de seguimiento.'),
    ('tratamientos', 'id, mascota_id, consulta_id', 'UUID', 'Mascota sí', 'PK; FK mascota; consulta opcional', 'Tratamiento actual o previo.'),
    ('servicios', 'id, nombre, activo', 'UUID, TEXT, BOOLEAN', 'Sí', 'PK', 'Catálogo de servicios.'),
    ('notificaciones', 'id, cliente_id, leida', 'UUID, BOOLEAN', 'Sí', 'PK; FK clientes', 'Mensajes del cliente.'),
], 7)
setp(135, '2.4.4 Contrato de la API de Supabase', 'Heading 2')
setp(136, 'Tabla 12. Operaciones utilizadas por Flutter y Edge Functions. Fuente: código del proyecto.', 'Normal')
setp(137, '')
fill_table(doc.tables[11], [
    ('Método', 'Ruta', 'Auth', 'Rol', 'Parámetros', '2xx esperado', 'Error / límite'),
    ('POST', '/auth/v1/token?grant_type=password', 'No', 'Cuenta existente', 'correo, contraseña', 'Sesión JWT', 'Credenciales inválidas'),
    ('POST', '/auth/v1/signup', 'No', 'Propietario', 'correo, contraseña', 'Usuario creado', 'Validación o correo duplicado'),
    ('POST', '/functions/v1/crear-propietario', 'JWT', 'Administrador', 'nombre, correo, contraseña, contacto', '201 + id', '400, 401, 403, 409, 500'),
    ('GET', '/rest/v1/clientes y /perfiles', 'JWT', 'Personal', 'select y filtros', 'Lista visible', 'RLS / permisos'),
    ('GET/POST/PATCH/DELETE', '/rest/v1/propietarios', 'JWT', 'Personal', 'id y JSON', 'Datos / 2xx', 'RLS / validación'),
    ('GET', '/rest/v1/mascotas', 'JWT', 'Personal; propietario propio', 'cliente_id, id, select', 'Filas visibles', 'RLS puede devolver 0 filas'),
    ('POST/PATCH/DELETE', '/rest/v1/mascotas', 'JWT', 'Administrador o veterinario', 'titular único; id; JSON', '2xx si autorizado', 'RLS o titular inválido'),
    ('GET/POST/PATCH/DELETE', '/rest/v1/consultas', 'JWT', 'Personal', 'mascota_id, id, JSON', 'Datos / 2xx', 'RLS / validación'),
    ('GET/POST/PATCH/DELETE', '/rest/v1/tratamientos y /vacunas', 'JWT', 'Personal; lectura autorizada móvil', 'mascota_id, id, JSON', 'Datos / 2xx', 'RLS / validación'),
    ('GET/POST/PATCH', '/rest/v1/citas', 'JWT', 'Personal; propietario propio', 'mascota_id, cliente_id, id, JSON', 'Datos / 2xx', 'RLS / mascota ajena'),
    ('GET', '/rest/v1/historiales_clinicos', 'JWT', 'Personal; propietario autorizado', 'mascota_id', 'Historial visible', 'RLS / autorizado=false'),
    ('POST/DELETE', '/storage/v1/object/avatars', 'JWT', 'Dueño del avatar', 'ruta y archivo', '2xx', 'Política Storage'),
    ('GET/HEAD', '/functions/v1/salud', 'No', 'Público', 'sin datos clínicos', '200 online', '503 offline; POST 200 observado: fallo'),
], 7)

# Stack y URLs exactas.
fill_table(doc.tables[12], [
    ('Componente', 'Versión', 'Función en el sistema', 'Criterio de selección'),
    ('Flutter / Dart', 'Flutter 3.44.8 / Dart 3.12.2 (caché local y netlify.toml)', 'Interfaces Android y web.', 'Una base de código para ambas vistas; RF-01 a RF-08.'),
    ('supabase_flutter', '^2.9.1 en pubspec.yaml; 2.17.2 resuelta en pubspec.lock', 'Auth, datos y Storage desde Flutter.', 'Integra servicios gestionados; RF-01 a RF-08.'),
    ('Supabase Auth', 'Servicio administrado, sin versión fijada', 'Identidad y sesiones JWT.', 'Control de acceso; RNF-02.'),
    ('PostgreSQL / PostgREST', 'Versiones remotas no confirmadas', 'Persistencia y API REST con RLS.', 'Relaciones y autorización; RNF-02.'),
    ('Netlify', 'Servicio de despliegue; Flutter 3.44.8 fijado', 'Web pública: https://eivetveterinaria.netlify.app/', 'Publicación desde GitHub; RNF-05.'),
    ('Supabase Edge Functions', 'Runtime remoto no confirmado', 'crear-propietario y salud.', 'Lógica segura de servidor; RF-02 y monitoreo.'),
    ('Git / GitHub Projects', 'Versiones cliente no necesarias para el usuario', 'Historial y tablero Kanban.', 'Trazabilidad de iteraciones.'),
], 7)
setp(139, 'Las versiones locales de Flutter y Dart provienen del archivo de caché del SDK; Netlify fija Flutter 3.44.8. La versión de los servicios administrados no se inventa. URL del repositorio: https://github.com/arturov7/Veterinaria-Eivet; URL del tablero: https://github.com/users/arturov7/projects/2/views/1?layout=board; ruta de salud: https://xbvjhrfgznljtwxghvka.supabase.co/functions/v1/salud.')
setp(140, 'Tabla 13. Stack tecnológico y despliegue real. Fuente: pubspec.yaml, netlify.toml y servicios publicados.')

# Pie de Figura 5; los wireframes son iniciales y no prueban permisos actuales.
fig5 = orig[134]
add_after(fig5, 'Figura 5. Wireframes iniciales de la aplicación móvil. La opción de registrar mascota representada allí fue retirada del rol propietario en la implementación actual. Fuente: elaboración propia.')

# Se elimina la lista editorial de verificación de la plantilla; los pendientes
# comprobables permanecen descritos en 2.7 y 2.8.
editorial = doc.tables[14]
editorial._element.getparent().remove(editorial._element)
orig[-1].text = 'Limitaciones documentadas: faltan mediciones del proceso actual de EIVET, acceso docente al tablero privado, verificación remota de RLS con dos propietarios ficticios y capturas fechadas de producción. La prueba automatizada de salud conserva un caso fallido (POST).'

# Bibliografía conforme al ejemplo Harvard / ISO 690-2 de la plantilla.
last = orig[-1]
last = add_after(last, 'BIBLIOGRAFÍA CONSULTADA', 'Heading 1')
references = [
    'DAYSMART VET. Cloud-Based Veterinary Practice Management Software [en línea]. s. f. Disponible en: <https://www.daysmart.com/vet/> [consulta: 03 octubre 2026].',
    'FLUTTER. Web support for Flutter [en línea]. s. f. Disponible en: <https://docs.flutter.dev/platform-integration/web> [consulta: 03 octubre 2026].',
    'KANBAN GUIDES. Open Guide to Kanban [en línea]. Versión 2025.7. 2025. Disponible en: <https://kanbanguides.org/open-guide-to-kanban/2025.7/pdf/open-guide-to-kanban.en-us.pdf> [consulta: 03 octubre 2026].',
    'SUPABASE. Auth [en línea]. s. f. Disponible en: <https://supabase.com/docs/guides/auth> [consulta: 03 octubre 2026].',
    'SUPABASE. Data REST API [en línea]. s. f. Disponible en: <https://supabase.com/docs/guides/api> [consulta: 03 octubre 2026].',
    'SUPABASE. Row Level Security [en línea]. s. f. Disponible en: <https://supabase.com/docs/guides/database/postgres/row-level-security> [consulta: 03 octubre 2026].',
    'VETCLOUD. Software de Gestión Integral para Clínicas y Hospitales Veterinarios [en línea]. s. f. Disponible en: <https://vetcloud.cl/> [consulta: 03 octubre 2026].',
    'VETICLOUD. Software de gestión para clínicas veterinarias [en línea]. s. f. Disponible en: <https://veticloud.com/es> [consulta: 03 octubre 2026].',
]
for item in references:
    last = add_after(last, item)

# Índices actualizados por número. La paginación debe actualizarse en Word.
cover_table = doc.tables[0]
intro_anchor = cover_table._element
index_lines = [
    ('ÍNDICE DE CONTENIDO', 'Heading 1'),
    ('Capítulo 1. Introducción · 1.1 Antecedentes · 1.2 Descripción del problema · 1.3 Planteamiento · 1.4 Justificación · 1.5 Objetivos', None),
    ('Capítulo 2. Marco teórico y desarrollo · 2.1 Fundamentos · 2.2 Metodología · 2.3 Requisitos · 2.4 Diseño · 2.5 Stack · 2.6 Implementación · 2.7 Seguridad · 2.8 Pruebas', None),
    ('ÍNDICE DE TABLAS', 'Heading 1'),
    ('Tablas 1 a 4: datos generales, soluciones, puntos críticos y objetivos. Tablas 5 a 9: iteraciones, actores, requisitos funcionales, no funcionales y alcance. Tablas 10 a 14: arquitectura, diccionario, API, stack y pruebas.', None),
    ('ÍNDICE DE FIGURAS', 'Heading 1'),
    ('Figura 1: tablero Kanban. Figura 2: casos de uso. Figura 3: arquitectura tecnológica. Figura 4: modelo de datos. Figura 5: wireframes iniciales.', None),
    ('La paginación de estos índices se actualizará en Word antes de exportar a PDF.', None),
]
for text, style in reversed(index_lines):
    p = doc.add_paragraph(text, style=style)
    intro_anchor.addnext(p._element)

caption5 = doc.add_paragraph('Tabla 5. Periodos de revisión Kanban y evidencias del repositorio. Fuente: elaboración propia.')
doc.tables[4]._element.addnext(caption5._element)

# Ajustes de portada / enlaces en la tabla general.
for row in doc.tables[0].rows:
    key = row.cells[0].text.strip()
    if key == 'Enlace al tablero de tareas':
        row.cells[1].text = 'https://github.com/users/arturov7/projects/2/views/1?layout=board (privado: 404 sin sesión; verificar permiso del docente).'
    elif key == 'Enlace al repositorio':
        row.cells[1].text = 'https://github.com/arturov7/Veterinaria-Eivet (acceso docente por verificar).'
    elif key == 'Organización o ámbito beneficiario':
        row.cells[1].text = 'Veterinaria EIVET, ciudad de Tarija, Bolivia.'
    elif key == 'Fecha de entrega':
        row.cells[1].text = '03/10/2026 · E3'

doc.save(OUTPUT)
print(OUTPUT.resolve())
