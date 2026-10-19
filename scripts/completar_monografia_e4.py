"""Crea el borrador E4 a partir del perfil E3, sin alterar el original."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


SOURCE = Path('Vedia_Salazar_Perfil_Proyecto_E3_corregido.docx')
OUTPUT = Path('Vedia_Salazar_Monografia_E4_borrador.docx')
doc = Document(SOURCE)
original = list(doc.paragraphs)


def before(anchor, text='', style=None):
    p = anchor.insert_paragraph_before(text, style=style)
    p.paragraph_format.space_after = Pt(6)
    return p


def section(anchor, title, paragraphs):
    before(anchor, title, 'Heading 2')
    for text in paragraphs:
        before(anchor, text)


def append(title, paragraphs):
    doc.add_paragraph(title, 'Heading 1')
    for text in paragraphs:
        doc.add_paragraph(text)


# Portada y datos generales: las comprobaciones sin evidencia quedan expresadas como límites.
original[3].text = 'MONOGRAFÍA FINAL · SISTEMA WEB Y MÓVIL PARA LA VETERINARIA EIVET'
original[4].text = 'Módulo 4 · Entregable E4 · Borrador de cierre · Tarija, octubre de 2026'
for row in doc.tables[0].rows[1:]:
    key = row.cells[0].text
    if key == 'Decisión de la tutoría T1':
        row.cells[1].text = 'Alcance documentado: gestión veterinaria; ventas, pagos e inventario avanzado quedan fuera del sistema.'
    elif key == 'Alcance acordado en T1':
        row.cells[1].text = 'Autenticación, propietarios, mascotas, citas, consultas, historial clínico, vacunas y tratamientos según RF-01 a RF-08.'
    elif key == 'Enlace al repositorio':
        row.cells[1].text = 'https://github.com/arturov7/Veterinaria-Eivet'
    elif key == 'Enlace al tablero de tareas':
        row.cells[1].text = 'https://github.com/users/arturov7/projects/2/views/1?layout=board (acceso privado; requiere invitación al evaluador).'
    elif key == 'Fecha de entrega':
        row.cells[1].text = 'E4: 10/10/2026 (plazo de la guía P4).'

# Sustituye índices de E3 y la instrucción de plantilla por preliminares utilizables.
for p in original[5:13]:
    p._element.getparent().remove(p._element)
anchor = original[13]
before(anchor, 'Cristian Arturo Vedia Salazar', 'Title')
before(anchor, 'Trabajo final presentado al Diplomado en Desarrollo Web y Aplicaciones Móviles de la Universidad Autónoma Juan Misael Saracho.')
before(anchor, 'Tarija, Bolivia · 2026')
before(anchor, 'HOJA DE APROBACIÓN', 'Heading 1')
before(anchor, 'Trabajo presentado a consideración del tribunal evaluador del Diplomado en Desarrollo Web y Aplicaciones Móviles.')
before(anchor, '______________________________     Tutor')
before(anchor, '______________________________     Miembro del tribunal 1')
before(anchor, '______________________________     Miembro del tribunal 2')
before(anchor, 'HOJA DE ADVERTENCIA', 'Heading 1')
before(anchor, 'El contenido, los resultados y las conclusiones de este trabajo son responsabilidad del autor. Su uso y reproducción deben reconocer la autoría y respetar los datos personales de la Veterinaria EIVET.')
before(anchor, 'RESUMEN', 'Heading 1')
before(anchor, 'La Veterinaria EIVET, en la ciudad de Tarija, necesita organizar la información de propietarios, mascotas y atención clínica en un sistema accesible para el personal y los propietarios. El objetivo fue desarrollar una aplicación web y móvil que centralice registros, consultas, citas e historial. Se aplicó Kanban para ordenar el trabajo en un tablero con límite de dos tareas en curso; la solución se implementó con Flutter, Supabase Auth, PostgreSQL con políticas de seguridad por fila y Netlify. El repositorio contiene los módulos de personal y propietario, las migraciones, diagramas y scripts de comprobación. En producción se observó el acceso al panel web y se midieron 30 solicitudes a la ruta de salud: todas respondieron HTTP 200 y 29 lo hicieron en tres segundos o menos. Esta medición no sustituye las pruebas autenticadas de las operaciones principales. La comprobación entre dos cuentas, el despliegue de la versión actual de algunas funciones y la validación con usuarios reales siguen pendientes. Se concluye que el proyecto dispone de una base funcional y documentada, pero su aceptación final requiere completar esas verificaciones y registrar evidencia reproducible.')
before(anchor, 'Palabras clave: veterinaria; Flutter; Supabase; Kanban; seguridad.')
before(anchor, 'ÍNDICE DE CONTENIDO', 'Heading 1')
for line in [
    'Capítulo 1. Introducción: antecedentes, problema, justificación y objetivos.',
    'Capítulo 2. Marco teórico y desarrollo: 2.1 Fundamentos; 2.2 Metodología; 2.3 Requisitos; 2.4 Diseño; 2.5 Stack; 2.6 Implementación; 2.7 Seguridad; 2.8 Pruebas; 2.9 Despliegue.',
    'Capítulo 3. Conclusiones y recomendaciones: 3.1 Conclusiones por objetivo; 3.2 Recomendaciones.',
    'Bibliografía consultada. Anexos A–F.'
]:
    before(anchor, line)
before(anchor, 'ÍNDICE DE TABLAS', 'Heading 1')
for line in [
    'Tablas 1–4. Datos generales, soluciones, puntos críticos y correspondencia de objetivos.',
    'Tablas 5–9. Iteraciones, actores, requisitos funcionales, requisitos no funcionales y alcance.',
    'Tablas 10–14. Arquitectura, diccionario de datos, contrato API, stack y casos de prueba.'
]:
    before(anchor, line)
before(anchor, 'ÍNDICE DE FIGURAS', 'Heading 1')
for line in [
    'Figura 1. Tablero Kanban de Veterinaria EIVET.',
    'Figura 2. Casos de uso y permisos.',
    'Figura 3. Componentes y protocolos de EIVET.',
    'Figura 4. Modelo de datos principal.',
    'Figura 5. Wireframes iniciales.'
]:
    before(anchor, line)
before(anchor, 'ÍNDICE DE ANEXOS', 'Heading 1')
before(anchor, 'Anexo A. Manual de usuario. Anexo B. Instalación y despliegue. Anexo C. Repositorio y tablero. Anexo D. Entregables ejecutables y URL. Anexo E. Diagramas. Anexo F. Evidencias y protocolo de pruebas.')

# 2.8: nueva evidencia con alcance explícito.
anchor = original[178]
original[177].text = ('Limitaciones al 03/10/2026: faltan mediciones del proceso manual de EIVET, acceso externo al tablero privado, verificación remota de RLS con dos propietarios ficticios, capturas fechadas de los flujos clínicos, prueba completa del CRUD y un APK final instalado. La medición de la ruta de salud figura en 2.9.1; sus resultados no prueban el rendimiento del CRUD.')
section(anchor, '2.9 Despliegue', [
    'La versión web se publica desde la rama principal del repositorio en Netlify. El archivo netlify.toml fija Flutter 3.44.8, ejecuta scripts/build_netlify.sh y publica build/web. El proceso obtiene dependencias, compila Flutter Web y entrega archivos estáticos por HTTPS. Las variables SUPABASE_URL y SUPABASE_PUBLISHABLE_KEY se suministran en Netlify; ninguna clave de servicio debe incluirse en el navegador o APK.',
    'Para reconstruir localmente se clona el repositorio, se instala Flutter 3.44.8, se ejecuta flutter pub get y se configura .env con los valores del proyecto Supabase. El panel se inicia con flutter run -d chrome --dart-define-from-file=.env. La aplicación Android se inicia seleccionando el identificador exacto de flutter devices; para generar un APK arm64 se usa flutter build apk --release --target-platform android-arm64 --dart-define-from-file=.env. El paquete resultante se encuentra en build/app/outputs/flutter-apk/app-release.apk. Estos pasos son instrucciones reproducibles; en esta revisión no se ejecutó una compilación nueva.',
    'La base de datos y los permisos se reconstruyen aplicando los scripts SQL del directorio supabase en el orden documentado en docs/SUPABASE_EIVET.md. La ruta de salud requiere la RPC de supabase/salud.sql y la Edge Function supabase/functions/salud; la función crear-propietario utiliza variables exclusivas del servidor. Los despliegues de funciones y migraciones requieren credenciales de administrador y deben verificarse en el proyecto remoto.',
    'La comprobación pública realizada el 03/10/2026 confirmó acceso al sitio web y 30 respuestas HTTP 200 de GET /functions/v1/salud. El script automatizado de métodos mostró GET y HEAD correctos, pero POST devolvió 200 en lugar del 405 esperado por el código local; por tanto, la versión desplegada de salud no se considera sincronizada. Tampoco se ha confirmado en producción la aplicación de las últimas políticas RLS, el CRUD completo, una APK final ni las pruebas con dos propietarios. El enlace del tablero requiere permiso para terceros.',
    'El procedimiento de verificación posterior al despliegue debe abrir el sitio, iniciar sesión con cuentas ficticias de cada rol, recorrer un alta y consulta autorizadas, comprobar una lectura y escritura cruzadas denegadas, probar la ruta de salud y registrar fecha, entorno, resultado y captura. Si falla un control, se revierte el cambio o se corrige la migración antes de aceptar el despliegue. El Anexo B detalla el procedimiento y el Anexo F enumera la evidencia disponible.'
])
section(anchor, '2.9.1 Medición de rendimiento en producción', [
    'El 03/10/2026 se ejecutó scripts/medir_rendimiento_salud.ps1 contra https://xbvjhrfgznljtwxghvka.supabase.co/functions/v1/salud. Se realizaron 30 solicitudes GET secuenciales con curl, pausa de 150 ms entre muestras y límite de 15 s por solicitud. Se registraron marca horaria, código HTTP y tiempo total percibido por el cliente en docs/evidencias/rendimiento_salud_2026-10-03.csv.',
    'Resultado observado: 30/30 respuestas HTTP 200; 29/30 (96,7 %) en 3 s o menos; promedio 0,896 s y percentil 90, 0,908 s. Esta muestra describe solo disponibilidad y latencia de la ruta de salud en un cliente y momento determinados. RNF-01 exige que al menos el 90 % de las operaciones principales responda en 3 s; ese requisito permanece sin aprobar porque faltan mediciones autenticadas de consultas, altas y modificaciones con usuarios de prueba.'
])
before(anchor, 'CAPÍTULO 3. CONCLUSIONES Y RECOMENDACIONES', 'Heading 1')
before(anchor, '3.1 Conclusiones por objetivo específico', 'Heading 2')
for item in [
    '1. Análisis: se identificaron actores, requisitos funcionales y no funcionales, alcance y flujo de trabajo (2.2 y 2.3). El levantamiento sigue siendo parcial porque la Tabla 3 aún contiene mediciones del proceso actual que deben tomarse en EIVET.',
    '2. Diseño: se definieron arquitectura, entidades, relaciones, permisos, contrato API y wireframes (2.4 y figuras 2–5). El diseño es coherente con el repositorio revisado, sujeto a comprobar el esquema y las políticas efectivamente aplicadas en el proyecto remoto.',
    '3. Implementación: el repositorio contiene panel web y aplicación móvil para propietarios, mascotas, citas, consultas, historial, vacunas y tratamientos (2.6). El CRUD y algunas rutas tienen código local, pero no se acreditó todavía su ejecución completa en producción.',
    '4. Seguridad: se integró Supabase Auth, validación de formularios, roles del personal y scripts RLS (2.7). No puede darse por satisfecha la separación entre propietarios hasta ejecutar y documentar pruebas negativas con dos cuentas en producción.',
    '5. Pruebas: se documentó una matriz con resultados esperados y observados, y se ejecutó una prueba automatizada de la ruta de salud (2.8). GET y HEAD cumplieron, mientras POST evidenció una discrepancia; la medición de salud no sustituye las pruebas de los flujos clínicos ni la verificación de RNF-01.',
    '6. Despliegue y documentación: la web pública responde en Netlify y se documentaron pasos de reconstrucción, despliegue y uso (2.9 y anexos). Una APK final, el despliegue sincronizado de la función de salud y la reproducción del manual por una segunda persona siguen por verificar.'
]:
    before(anchor, item)
section(anchor, '3.2 Recomendaciones', [
    'Aplicar y auditar las migraciones RLS pendientes; crear dos cuentas ficticias de propietario y conservar evidencia de que cada una solo puede leer sus registros y no puede modificarlos indebidamente. Repetir los casos con administrador y veterinario para verificar los permisos previstos.',
    'Publicar la versión de la función salud que rechaza POST, ejecutar de nuevo scripts/verify_salud.ps1 y adjuntar salida y captura fechada. Completar las pruebas funcionales del CRUD web y móvil y obtener una APK arm64 instalable.',
    'Medir RNF-01 en producción sobre operaciones principales con sesiones autorizadas, una muestra y condiciones definidas, guardando tiempos individuales y criterio de aceptación. Cronometrar también el proceso manual actual de EIVET para reemplazar los valores no medidos de la Tabla 3.',
    'Solicitar acceso de lectura al tablero privado para el tribunal, confirmar los datos institucionales de portada y ejecutar una revisión del manual con otra persona. Ensayar una demostración de siete minutos con cuentas ficticias, red estable, ruta de salud y evidencias ya preparadas.'
])

append('ANEXO A. MANUAL DE USUARIO', [
    'A.1 Acceso. Abra https://eivetveterinaria.netlify.app/ desde un navegador actualizado. El personal ingresa con una cuenta habilitada por administración. La aplicación móvil Android se instala desde el APK generado para el dispositivo; los propietarios usan sus credenciales de Supabase Auth. Nunca utilice datos clínicos reales en demostraciones.',
    'A.2 Personal administrador y veterinario. Tras iniciar sesión, use el menú lateral para localizar propietarios, mascotas, citas veterinarias, consultas e historial, vacunas y tratamientos. Seleccione un registro antes de editarlo; compruebe titular, mascota y fecha antes de guardar. El administrador puede crear cuentas de propietario mediante el flujo autorizado. La gestión de mascotas corresponde al administrador y al veterinario.',
    'A.3 Propietario. En la app móvil, abra Mascotas para consultar la ficha propia; use Citas para solicitar o consultar una cita y las pestañas de la tarjeta para ver historial, vacunas y tratamientos disponibles. El propietario no debe registrar, editar ni eliminar mascotas. Si una sección carece de registros, se presenta un estado vacío.',
    'A.4 Errores y ayuda. Si el inicio de sesión falla, revise correo y contraseña o solicite restablecimiento al administrador. Si no aparece una mascota, confirme la cuenta y comunique al personal su identificador; no intente acceder con la cuenta de otra persona. Ante un error de red, reintente cuando haya conexión y registre hora y operación.'
])
append('ANEXO B. MANUAL DE INSTALACIÓN Y DESPLIEGUE', [
    'B.1 Requisitos. Instale Git y Flutter 3.44.8 con Dart 3.12.2; para Android agregue Android SDK y un equipo autorizado por USB. Se requiere un proyecto Supabase y acceso de despliegue a Netlify. Las credenciales sensibles se administran fuera del repositorio.',
    'B.2 Preparación local. Ejecute git clone https://github.com/arturov7/Veterinaria-Eivet.git, ingrese al directorio del proyecto y ejecute flutter pub get. Copie .env.example a .env y complete APP_MODE, SUPABASE_URL y SUPABASE_PUBLISHABLE_KEY con los datos de su entorno. No publique .env.',
    'B.3 Base de datos y funciones. Aplique los SQL indicados en docs/SUPABASE_EIVET.md en el orden allí descrito. Verifique tablas, índices, RLS y políticas con cuentas ficticias antes de cargar información. Instale la RPC de supabase/salud.sql; despliegue las Edge Functions del directorio supabase/functions y configure sus secretos de servidor. Vuelva a ejecutar las comprobaciones de salud y de aislamiento de cuentas.',
    'B.4 Ejecución y compilación. Para web local: flutter run -d chrome --dart-define-from-file=.env. Para Android: flutter devices y flutter run -d IDENTIFICADOR_REAL --dart-define-from-file=.env; reemplace IDENTIFICADOR_REAL por la columna ID, sin escribir signos de menor o mayor. Para APK arm64: flutter build apk --release --target-platform android-arm64 --dart-define-from-file=.env. Para web: flutter build web --release --dart-define-from-file=.env. Estos comandos se documentan para reproducir el despliegue; no se ejecutaron en esta revisión.',
    'B.5 Netlify. Conecte la rama main de GitHub; netlify.toml ejecuta scripts/build_netlify.sh y publica build/web. En la configuración del sitio cargue SUPABASE_URL y SUPABASE_PUBLISHABLE_KEY. Tras desplegar, abra la URL pública, pruebe inicio de sesión con cuentas ficticias y registre identificador de despliegue, fecha, capturas y resultados de pruebas.',
    'B.6 Comprobación por otra persona. Entregue este manual y acceso temporal a un entorno de prueba a una persona distinta del autor. Pídale registrar sistema operativo, versión Flutter, comandos ejecutados, resultado de la compilación y una captura de inicio de sesión. Esta prueba de reproducibilidad aún no se ha realizado.'
])
append('ANEXO C. REPOSITORIO Y TABLERO', [
    'Repositorio: https://github.com/arturov7/Veterinaria-Eivet. Rama de publicación: main. El código contiene lib, scripts, supabase, netlify.toml, pubspec.yaml y documentación técnica. El README ofrece una entrada rápida.',
    'Tablero Kanban: https://github.com/users/arturov7/projects/2/views/1?layout=board. La Figura 1 conserva la captura del tablero; su acceso externo es privado y debe compartirse con el evaluador. El límite de trabajo en curso es de dos tarjetas.'
])
append('ANEXO D. URL Y PAQUETE INSTALABLE', [
    'Versión web pública: https://eivetveterinaria.netlify.app/. Ruta de salud: https://xbvjhrfgznljtwxghvka.supabase.co/functions/v1/salud. La apertura pública del panel y 30 respuestas GET de salud se comprobaron el 03/10/2026.',
    'El APK arm64 se obtiene con el comando del Anexo B en build/app/outputs/flutter-apk/app-release.apk. No se adjunta un APK verificado en este borrador; debe compilarse, instalarse y registrar versión, huella y evidencia de uso antes de la entrega.'
])
append('ANEXO E. DIAGRAMAS Y TRAZABILIDAD', [
    'Las Figuras 2–5 del cuerpo del documento muestran casos de uso, componentes y protocolos, modelo de datos y wireframes. La Tabla 4 relaciona los seis objetivos con evidencias y conclusiones; las Tablas 10–12 detallan arquitectura, diccionario y contrato API. Esos elementos integran el anexo técnico y deben contrastarse con el proyecto Supabase desplegado antes del cierre.'
])
append('ANEXO F. EVIDENCIAS, CASOS DE PRUEBA Y DEFENSA', [
    'F.1 Evidencia disponible. La Figura 1 documenta el tablero Kanban. docs/PRUEBAS_Y_EVIDENCIAS.md contiene la matriz de casos; scripts/verify_salud.ps1 automatiza métodos de la ruta de salud; docs/evidencias/rendimiento_salud_2026-10-03.csv conserva las 30 muestras de rendimiento y scripts/medir_rendimiento_salud.ps1 permite repetirlas.',
    'F.2 Resultado observado. En salud, GET y HEAD fueron aprobados; POST falló porque producción devolvió HTTP 200 y el código local espera 405. La prueba con dos propietarios, CRUD completo, capturas fechadas de los flujos clínicos y compilación/instalación del APK permanecen sin evidencia de aprobación.',
    'F.3 Ensayo de defensa. Preparar una demostración de siete minutos: problema y objetivo (1 min), arquitectura y seguridad (1 min), inicio de sesión y flujo principal en producción (3 min), caso de permiso denegado y salud (1 min), resultados y límites (1 min). Reservar cinco minutos para preguntas. Cuentas ficticias, red de respaldo, capturas y CSV deben estar disponibles antes del ensayo. El ensayo con una segunda persona todavía no se ha efectuado.'
])

# Conservar formato base y evitar espacios excesivos en los nuevos párrafos.
for p in doc.paragraphs:
    if p.text.startswith('2.9 Despliegue'):
        p.style = 'Heading 1'
    elif p.text.startswith('2.9.1 Medición'):
        p.style = 'Heading 2'
    elif p.text.startswith(('CU-01 Iniciar sesión', 'CU-02 Registrar mascota', 'CU-03 Registrar consulta')):
        p.style = 'Heading 3'
    elif p.text.startswith(('El proceso de registro y seguimiento', 'Desarrollar una aplicación web y móvil multiplataforma', 'Arquitectura cliente')) or (not p.text.strip() and p.style.name.startswith('Heading')):
        p.style = 'Normal'
    if p.style.name.startswith('Heading'):
        p.paragraph_format.keep_with_next = True
    if p.text.startswith(('HOJA DE APROBACIÓN', 'HOJA DE ADVERTENCIA', 'RESUMEN', 'ÍNDICE DE CONTENIDO', 'CAPÍTULO ', 'BIBLIOGRAFÍA CONSULTADA', 'ANEXO A.')):
        p.paragraph_format.page_break_before = True
doc.save(OUTPUT)
print(OUTPUT.resolve())
