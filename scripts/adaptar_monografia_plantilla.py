"""Compone la monografía E4 sobre la plantilla institucional de Word."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


TEMPLATE = Path(r'C:\Users\Arturo Vedia\Downloads\Plantilla_Monografia_Diplomado.docx')
SOURCE = Path('Vedia_Salazar_Monografia_E4_borrador.docx')
OUTPUT = Path('Vedia_Salazar_Monografia_E4_adaptada.docx')
template = Document(TEMPLATE)
source = Document(SOURCE)
body = template.element.body
break_el = template.paragraphs[47]._p  # salto institucional, numeración romana -> arábiga
final_section = body.sectPr
samples = {name: deepcopy(template.paragraphs[index]._p) for name, index in {
    'institution': 0, 'title': 4, 'coverline': 6, 'right': 8,
    'major': 35, 'minor': 52, 'body': 45,
}.items()}

# La plantilla permanece como paquete base. Se conserva el salto de sección y el pie.
for child in list(body):
    if child is not break_el and child is not final_section:
        body.remove(child)


def add_prelim(text, kind='body', page=False):
    el = deepcopy(samples[kind])
    for child in list(el):
        if child.tag != qn('w:pPr'):
            el.remove(child)
    break_el.addprevious(el)
    p = Paragraph(el, template._body)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt({'institution': 13, 'title': 16, 'coverline': 12, 'right': 12,
                        'major': 13, 'minor': 11, 'body': 12}[kind])
    run.bold = kind in ('institution', 'title', 'major')
    if page:
        p.paragraph_format.page_break_before = True
    return p


title = 'Sistema web y móvil para personal veterinario y propietarios de la Veterinaria EIVET en Tarija'
for line in [
    'UNIVERSIDAD AUTÓNOMA “JUAN MISAEL SARACHO”',
    'FACULTAD DE CIENCIAS Y TECNOLOGÍA',
    'DIRECCIÓN DE POSGRADO',
    'CARRERA DE INGENIERÍA INFORMÁTICA',
]:
    add_prelim(line, 'institution')
add_prelim(title.upper(), 'title')
add_prelim('Trabajo final presentado para la obtención del título de', 'coverline')
add_prelim('DIPLOMADO EN DESARROLLO WEB Y APLICACIONES MÓVILES', 'institution')
add_prelim('Postulante: Cristian Arturo Vedia Salazar', 'right')
add_prelim('Tutor: [NOMBRE DEL TUTOR POR CONFIRMAR]', 'right')
add_prelim('Tarija – Bolivia', 'coverline')
add_prelim('2026', 'coverline')

add_prelim('CONTRATAPA', 'major', page=True)
add_prelim(title.upper(), 'title')
add_prelim('Postulante: Cristian Arturo Vedia Salazar', 'right')
add_prelim('Tutor: [NOMBRE DEL TUTOR POR CONFIRMAR]', 'right')
add_prelim('Autoridades y tribunal: [NÓMINA INSTITUCIONAL POR CONFIRMAR]', 'body')
add_prelim('Tarija – Bolivia · 2026', 'coverline')

add_prelim('HOJA DE APROBACIÓN', 'major', page=True)
add_prelim('Trabajo final presentado a consideración del tribunal evaluador del Diplomado en Desarrollo Web y Aplicaciones Móviles.')
add_prelim('__________________________________    Tutor [FIRMA Y NOMBRE]', 'body')
add_prelim('__________________________________    Tribunal 1 [FIRMA Y NOMBRE]', 'body')
add_prelim('__________________________________    Tribunal 2 [FIRMA Y NOMBRE]', 'body')
add_prelim('HOJA DE ADVERTENCIA', 'major', page=True)
add_prelim('[TEXTO DE ADVERTENCIA INSTITUCIONAL POR CONFIRMAR]', 'body')

add_prelim('ÍNDICE DE CONTENIDO', 'major', page=True)
for line in [
    'Capítulo 1. Introducción — antecedentes, problema, justificación y objetivos.',
    'Capítulo 2. Marco teórico — fundamentos, Kanban, requisitos, diseño, stack, implementación, seguridad, pruebas y despliegue.',
    'Capítulo 3. Conclusiones y recomendaciones.',
    'Bibliografía consultada. Anexos A–F.',
]:
    add_prelim(line)
add_prelim('ÍNDICE DE TABLAS', 'major', page=True)
table_titles = [
    'Soluciones informáticas existentes', 'Referencias externas para puntos críticos',
    'Correspondencia entre objetivos, evidencia y conclusiones',
    'Periodos de revisión Kanban', 'Actores del sistema',
    'Requisitos funcionales', 'Requisitos no funcionales',
    'Capacidades fuera de alcance', 'Componentes de la arquitectura',
    'Diccionario de datos', 'Contrato de la API',
    'Stack tecnológico', 'Casos de prueba y resultados',
]
for i, caption in enumerate(table_titles, 1):
    add_prelim(f'Tabla {i}. {caption}.')
add_prelim('ÍNDICE DE FIGURAS', 'major', page=True)
for i, caption in enumerate([
    'Tablero Kanban de EIVET', 'Casos de uso y permisos',
    'Componentes y protocolos', 'Modelo de datos', 'Wireframes iniciales',
], 1):
    add_prelim(f'Figura {i}. {caption}.')
add_prelim('ÍNDICE DE ANEXOS', 'major', page=True)
for letter, caption in zip('ABCDEF', [
    'Manual de usuario', 'Manual de instalación y despliegue',
    'Enlace al repositorio', 'Dirección pública o instalable',
    'Diagramas', 'Evidencias y documentación complementaria',
]):
    add_prelim(f'Anexo {letter}. {caption}.')

add_prelim('RESUMEN', 'major', page=True)
summary = (
    'La Veterinaria EIVET, en Tarija, requiere organizar información de propietarios, mascotas y atención clínica. '
    'Este trabajo desarrolla un sistema web para el personal veterinario y una aplicación móvil de consulta y solicitud '
    'de citas para propietarios. El proyecto utilizó Kanban para ordenar tareas y Flutter, Supabase y Netlify para la '
    'implementación. El repositorio contiene interfaces, modelo de datos, políticas de seguridad y scripts de comprobación. '
    'Se observó el acceso a la web pública y se ejecutó una prueba automatizada de la ruta de salud: los métodos GET y HEAD '
    'cumplieron, mientras que POST devolvió 200 en vez del 405 esperado. Treinta solicitudes GET a esa ruta respondieron 200; '
    '29 tardaron tres segundos o menos. Esta medición no demuestra el rendimiento de las operaciones clínicas. La separación '
    'entre dos propietarios, varios flujos completos en producción y el instalable final aún requieren comprobación. El '
    'sistema cuenta con una base implementada y documentada, cuya aceptación depende de completar esas verificaciones.'
)
add_prelim(summary)
add_prelim('Palabras clave: veterinaria; aplicación móvil; Flutter; Supabase; Kanban.')

# Traslada el contenido desde el Capítulo 1. Omite la tabla preliminar de perfil.
start = next(p._p for p in source.paragraphs if p.text.startswith('CAPÍTULO 1. INTRODUCCIÓN'))
copying = False
for child in list(source.element.body):
    if child is start:
        copying = True
    if not copying or child.tag == qn('w:sectPr'):
        continue
    clone = deepcopy(child)
    for pstyle in clone.xpath('.//w:pPr/w:pStyle'):
        mapping = {'Ttulo1': 'Heading1', 'Ttulo2': 'Heading2', 'Ttulo3': 'Heading3'}
        old_style = pstyle.get(qn('w:val'))
        if old_style in mapping:
            pstyle.set(qn('w:val'), mapping[old_style])
        elif old_style in ('NormalWeb', 'Prrafodelista', 'pdq2pgselectionanchorcontainer'):
            pstyle.getparent().remove(pstyle)
    # Los cinco PNG verificables se reinsertan como relaciones del documento nuevo.
    for blip in clone.xpath('.//a:blip'):
        old = blip.get(qn('r:embed'))
        if old:
            part = source.part.related_parts[old]
            new, _ = template.part.get_or_add_image(BytesIO(part.blob))
            blip.set(qn('r:embed'), new)
    for img_layer in clone.xpath('.//w14:imgLayer'):
        img_layer.getparent().remove(img_layer)
    final_section.addprevious(clone)


def replace_text(p, value):
    if p.text != value:
        p.text = value


def renumber(text):
    def change(match):
        first = int(match.group(2))
        second = match.group(3)
        if first == 1:
            return match.group(0)
        return f'{match.group(1)} {first-1}' + (f'–{int(second)-1}' if second else '')
    return re.sub(r'\b(Tabla|Tablas)\s+(\d+)(?:[–-](\d+))?', change, text)


main = list(template.paragraphs)[len(template.paragraphs)-len(source.paragraphs)+36:] if False else template.paragraphs
in_main = False
for p in template.paragraphs:
    if p.text.startswith('CAPÍTULO 1. INTRODUCCIÓN'):
        in_main = True
    if not in_main:
        continue
    text = p.text
    updated = renumber(text)
    updated = updated.replace('CAPÍTULO 2. MARCO TEÓRICO Y DESARROLLO', 'CAPÍTULO 2. MARCO TEÓRICO')
    updated = updated.replace('2.3.4 Requisitos declarados fuera de alcance', 'Alcance y exclusiones')
    updated = updated.replace('2.3.5 Casos de uso', '2.3.4 Casos de uso')
    updated = updated.replace('2.4.4 Contrato de la API de Supabase', '2.4.4 Diseño de la interfaz de programación de aplicaciones')
    updated = updated.replace('3.1 Conclusiones por objetivo específico', '3.1 Conclusiones')
    if text.startswith('Actualmente, la veterinaria realiza sus procesos de atención utilizando métodos manuales'):
        updated = ('El diagnóstico inicial del proyecto describe procedimientos manuales para registrar propietarios, '
                   'mascotas y atenciones en EIVET. La magnitud de los tiempos, errores y duplicaciones aún requiere '
                   'observación y medición directa en la clínica. La Tabla 2 presenta resultados publicados de otros '
                   'establecimientos veterinarios como contexto, sin atribuirlos a EIVET.')
    if updated.startswith('Tabla 2. Puntos críticos del proceso actual.'):
        updated = 'Tabla 2. Puntos críticos y datos publicados de otras veterinarias. Fuente: Ahmed et al. (2023) y Jones-Diette et al. (2017); no son mediciones de EIVET.'
    if text.startswith('1. Análisis: se identificaron actores'):
        updated = updated.replace('la Tabla 2 aún contiene mediciones del proceso actual que deben tomarse en EIVET',
                                  'la Tabla 2 contiene referencias externas y todavía faltan mediciones propias de EIVET')
    if text.startswith('Figura 5. Wireframes'):
        updated = ('Figura 5. Wireframes iniciales de la aplicación móvil. La opción de registrar mascotas para el propietario '
                   'es un diseño previo y no corresponde a los permisos vigentes. Fuente: elaboración propia.')
    if text.startswith('Anexo') or text.startswith('ANEXO'):
        if re.match(r'ANEXO [A-F]\.', text):
            updated = text.replace('ANEXO ', 'Anexo ', 1)
            p.style = 'Heading2'
    if updated != text:
        replace_text(p, updated)
    if p.text.startswith('CAPÍTULO ') and not p.text.startswith('CAPÍTULO 1.'):
        p.paragraph_format.page_break_before = True
    if p.text.startswith('BIBLIOGRAFÍA CONSULTADA'):
        p.paragraph_format.page_break_before = True
    if re.match(r'^(Tabla|Figura) \d+\.', p.text):
        p.paragraph_format.keep_with_next = p.text.startswith('Tabla')
    if p._p.xpath('.//a:blip'):
        p.paragraph_format.keep_with_next = True
    if re.match(r'^Anexo [A-F]\.', p.text):
        p.paragraph_format.page_break_before = True
    if p.style and p.style.name.startswith('Heading'):
        p.paragraph_format.keep_with_next = True
    elif text and not p._p.xpath('.//a:blip'):
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if text.startswith(('Tabla ', 'Figura ')) else WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.0 if text.startswith(('Tabla ', 'Figura ')) else 1.5
        p.paragraph_format.space_after = Pt(8.5)
    for run in p.runs:
        if not run.text:
            continue
        run.font.name = 'Arial'
        if p.style and p.style.name == 'Heading 1':
            run.font.size = Pt(13)
            run.bold = True
        elif p.style and p.style.name in ('Heading 2', 'Heading 3'):
            run.font.size = Pt(11)
            run.bold = True
        else:
            run.font.size = Pt(12)

# Inserta el título general de anexos, previsto en la plantilla.
annex = next(p for p in template.paragraphs if p.text.startswith('Anexo A. MANUAL'))
heading = annex.insert_paragraph_before('ANEXOS', style='Heading1')
heading.paragraph_format.page_break_before = True
annex.paragraph_format.page_break_before = False

# El Anexo E reproduce los diagramas presentes en el borrador, a ancho legible.
annex_f = next(p for p in template.paragraphs if p.text.startswith('Anexo F. EVIDENCIAS'))
for image_number, label in [
    (1, 'E.1 Casos de uso y permisos.'),
    (2, 'E.2 Componentes y protocolos.'),
    (3, 'E.3 Modelo de datos.'),
    (4, 'E.4 Wireframes iniciales; el botón de registro de mascotas del propietario corresponde a un diseño anterior.'),
]:
    title_p = annex_f.insert_paragraph_before(label)
    title_p.paragraph_format.keep_with_next = True
    picture_p = annex_f.insert_paragraph_before()
    picture_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    old_shape = source.inline_shapes[image_number]
    old_rid = old_shape._inline.xpath('.//a:blip')[0].get(qn('r:embed'))
    old_part = source.part.related_parts[old_rid]
    picture_p.add_run().add_picture(BytesIO(old_part.blob), width=Cm(14.3))

# Dos referencias académicas comprobadas complementan la documentación oficial.
antecedents = next(p for p in template.paragraphs if p.text == '1.1 Antecedentes')
antecedent_note = antecedents.insert_paragraph_before(
    'Los registros clínicos electrónicos veterinarios permiten recuperar y analizar información de la atención. '
    'Anholt et al. (2014) documentaron la extracción de registros de consultorios para su uso en vigilancia y '
    'estudios basados en evidencia. Ese antecedente respalda la necesidad de estructurar la información clínica, '
    'aunque no aporta mediciones del proceso particular de EIVET.'
)
antecedents._p.addnext(antecedent_note._p)
fundamentals = next(p for p in template.paragraphs if p.text.startswith('2.1 Fundamentos conceptuales'))
fundamentals_note = fundamentals.insert_paragraph_before(
    'En la comunicación entre componentes se adoptan solicitudes HTTP y datos JSON. Fielding (2000) expone '
    'los principios arquitectónicos de los sistemas distribuidos en red; en EIVET, las rutas concretas y sus '
    'permisos se contrastan con el código y las políticas del proyecto, no se deducen solo del modelo teórico.'
)
fundamentals._p.addnext(fundamentals_note._p)
refs = next(p for p in template.paragraphs if p.text.startswith('DAYSMART VET.'))
refs.insert_paragraph_before(
    'AHMED, M. A.; OTSYINA, H. R.; JOHNSON, S. A.; ASARE, D. A.; SHABAN, S.; EMIKPE, B. O. '
    'User perception and attitude toward electronic medical record system in veterinary practice in Ghana [en línea]. '
    'PAMJ-One Health, 12, artículo 14, 2023. DOI: 10.11604/pamj-oh.2023.12.14.40720. Disponible en: '
    '<https://www.one-health.panafrican-med-journal.com/content/article/12/14/full/> [consulta: 03 octubre 2026].'
)
refs.insert_paragraph_before(
    'ANHOLT, R. M.; BEREZOWSKI, J.; MACLEAN, K.; RUSSELL, M. L.; JAMAL, I.; STEPHEN, C. '
    'The application of medical informatics to the veterinary management programs at companion animal practices '
    'in Alberta, Canada: a case study. Preventive Veterinary Medicine [en línea], 113(2), 165–174, 2014. '
    'DOI: 10.1016/j.prevetmed.2013.11.005. Disponible en: <https://pubmed.ncbi.nlm.nih.gov/24299904/> '
    '[consulta: 03 octubre 2026].'
)
flutter_ref = next(p for p in template.paragraphs if p.text.startswith('FLUTTER.'))
flutter_ref.insert_paragraph_before(
    'FIELDING, R. T. Architectural Styles and the Design of Network-based Software Architectures [en línea]. '
    'Tesis doctoral. University of California, Irvine, 2000. Disponible en: '
    '<https://ics.uci.edu/~fielding/pubs/dissertation/abstract.htm> [consulta: 03 octubre 2026].'
)
kanban_ref = next(p for p in template.paragraphs if p.text.startswith('KANBAN GUIDES.'))
kanban_ref.insert_paragraph_before(
    'JONES-DIETTE, J.; ROBINSON, N. J.; COBB, M.; BRENNAN, M. L.; DEAN, R. S. '
    'Accuracy of the electronic patient record in a first opinion veterinary practice [en línea]. '
    'Preventive Veterinary Medicine, 148, 121–126, 2017. DOI: 10.1016/j.prevetmed.2016.11.014. '
    'Disponible en: <https://doi.org/10.1016/j.prevetmed.2016.11.014> '
    '[consulta: 03 octubre 2026].'
)

# Ajusta las tablas al ancho útil de la hoja carta con márgenes institucionales.
critical = template.tables[1]
critical.rows[0].cells[1].text = 'Aspecto comparable'
critical.rows[0].cells[2].text = 'Resultado publicado (no EIVET)'
critical.rows[1].cells[1].text = 'Tiempo para recuperar información del paciente'
critical.rows[1].cells[2].text = 'En 7 centros observados de Ghana, 6 tardaron 0–20 min y 1 tardó 21–40 min (Ahmed et al., 2023).'
critical.rows[2].cells[1].text = 'Tiempo para registrar información del paciente'
critical.rows[2].cells[2].text = 'Los 7 centros observados emplearon 21–40 min para registrar datos del paciente (Ahmed et al., 2023).'
critical.rows[3].cells[1].text = 'Completitud de los registros clínicos'
critical.rows[3].cells[2].text = 'En 36 consultas de una práctica británica se registró el 64,4 % de los problemas tratados y el 58,3 % de las acciones observadas (Jones-Diette et al., 2017). No mide duplicados.'
available = template.sections[-1].page_width - template.sections[-1].left_margin - template.sections[-1].right_margin
for shape in template.inline_shapes:
    if shape.width > available - Cm(0.2):
        ratio = (available - Cm(0.2)) / shape.width
        shape.width = int(shape.width * ratio)
        shape.height = int(shape.height * ratio)
for table in template.tables:
    table.autofit = False
    borders = table._tbl.tblPr.find(qn('w:tblBorders'))
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        table._tbl.tblPr.append(borders)
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        item = borders.find(qn(f'w:{side}'))
        if item is None:
            item = OxmlElement(f'w:{side}')
            borders.append(item)
        item.set(qn('w:val'), 'single')
        item.set(qn('w:sz'), '4')
        item.set(qn('w:color'), 'D9D9D9')
    widths = [col.width for col in table.columns]
    total = sum(widths)
    if total > available:
        scale = available / total
        new_widths = [int(width * scale) for width in widths]
        for col, width in zip(table.columns, new_widths):
            col.width = width
        for row in table.rows:
            for cell, width in zip(row.cells, new_widths):
                cell.width = width
    for row_number, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_number == 0:
                shading = cell._tc.get_or_add_tcPr().find(qn('w:shd'))
                if shading is None:
                    shading = OxmlElement('w:shd')
                    cell._tc.get_or_add_tcPr().append(shading)
                shading.set(qn('w:fill'), 'EDEDED')
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = 'Arial'
                    run.font.size = Pt(8 if len(table.columns) >= 6 else 9)
                    if row_number == 0:
                        run.bold = True
        if row_number == 0:
            trPr = row._tr.get_or_add_trPr()
            repeat = OxmlElement('w:tblHeader')
            repeat.set(qn('w:val'), 'true')
            trPr.append(repeat)

# La plantilla solicita refrescar campos. No se afirma que los índices tengan páginas calculadas.
settings = template.settings.element
update = settings.find(qn('w:updateFields'))
if update is None:
    update = OxmlElement('w:updateFields')
    settings.append(update)
update.set(qn('w:val'), 'true')

for paragraph in template.paragraphs:
    for run in paragraph.runs:
        if run.text and run.font.size is None:
            run.font.name = 'Arial'
            run.font.size = Pt(12)

template.save(OUTPUT)
print(OUTPUT.resolve())
