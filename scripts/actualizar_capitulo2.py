"""Añade 2.6–2.8 al perfil sin sobrescribir el archivo de origen."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt

SOURCE = Path('Vedia_Salazar_Perfil_Proyecto_v2_corregido.docx')
OUTPUT = Path('Vedia_Salazar_Avance_Capitulo2_Implementacion_Seguridad_Pruebas.docx')
doc = Document(SOURCE)
anchor = doc.tables[13]._element  # Lista de control editorial, posterior al capítulo.


def paragraph(text, style=None):
    p = doc.add_paragraph(text, style=style)
    anchor.addprevious(p._element)
    return p


def heading(text, level=2):
    return paragraph(text, f'Heading {level}')


heading('2.6 Implementación', 1)
paragraph('La solución se implementó en Flutter con una entrada común. En la versión web se muestra el panel para el personal y en Android la aplicación del propietario; la selección ocurre en lib/main.dart mediante kIsWeb. La configuración de Supabase se recibe en tiempo de compilación con APP_MODE, SUPABASE_URL y SUPABASE_PUBLISHABLE_KEY. Los repositorios separan las pantallas de las consultas a Supabase y permiten un modo demo para desarrollo. El modo demo no constituye evidencia de funcionamiento de producción.')
paragraph('La aplicación móvil incluye inicio de sesión, consulta de mascotas asociadas a la cuenta, solicitudes de citas y visualización de cuidados. El panel web reúne la gestión de propietarios, mascotas, citas, consultas, vacunas y tratamientos. Supabase Auth gestiona las identidades; las tablas de PostgreSQL almacenan los registros clínicos. La función crear-propietario reserva la creación de cuentas desde el panel a un administrador autenticado y la función salud consulta una RPC de solo lectura. La publicación web se realiza mediante Netlify desde el repositorio GitHub.')
paragraph('Una dificultad real apareció en el formulario móvil de mascotas: valores históricos de sexo con distinta capitalización causaban una aserción del desplegable de Flutter. Se resolvió normalizando el valor antes de asignarlo al control y dejando dos opciones únicas, Macho y Hembra. El cambio consta en el commit 7dd5379. Esta corrección de interfaz no demuestra por sí misma que las restricciones de acceso a mascotas estén activas en el proyecto remoto.')

heading('2.7 Seguridad', 1)
heading('2.7.1 Autenticación y gestión de sesiones', 2)
paragraph('La aplicación usa Supabase Auth para iniciar y cerrar sesión, registrar cuentas móviles y recuperar contraseñas. El cliente observa los cambios de sesión y asocia las preferencias locales con el identificador del usuario. La existencia de estas llamadas se verificó en el código; no se documenta aquí una prueba de caducidad o renovación de sesión en producción.')
heading('2.7.2 Autorización por rol', 2)
paragraph('El esquema define los roles administrador y veterinario en public.perfiles. La función crear-propietario valida el JWT y consulta el rol administrador antes de crear una cuenta. La política prevista para mascotas permite al propietario leer las vinculadas con su cuenta y reserva el alta, edición y eliminación al personal. La política prevista para citas también comprueba el vínculo con la mascota. Los scripts 12_MASCOTAS_TITULAR_SOLO_PERSONAL.sql y 13_CITAS_RLS_PROPIETARIOS.sql describen estas reglas, pero su aplicación efectiva y el aislamiento entre dos cuentas siguen pendientes de comprobación en Supabase. Por ello no se afirma todavía que la protección RLS esté aprobada en producción.')
heading('2.7.3 Validación de entradas y vulnerabilidades habituales', 2)
paragraph('Los formularios móviles comprueban campos obligatorios antes de guardar. La función crear-propietario rechaza métodos distintos de POST, JSON inválido, correo con formato incorrecto y contraseñas de menos de ocho caracteres; también exige una sesión válida y rol administrador. Las políticas RLS deben impedir el acceso directo no autorizado por PostgREST, incluso si alguien evita la interfaz. Esa defensa no se considera verificada hasta ejecutar los casos de dos cuentas. La ruta salud solo devuelve online u offline y no lee datos clínicos.')
heading('2.7.4 Cifrado de credenciales y comunicaciones', 2)
paragraph('Las contraseñas se entregan a Supabase Auth y no se guardan en las tablas de la aplicación. Las URL públicas del panel y de Supabase emplean HTTPS; se comprobó su acceso por esa vía. No se realizó una auditoría independiente de los mecanismos internos de cifrado del proveedor ni una inspección de todo el tráfico móvil.')
heading('2.7.5 Gestión de variables sensibles', 2)
paragraph('El archivo .env local está excluido del repositorio; Netlify recibe SUPABASE_URL y SUPABASE_PUBLISHABLE_KEY como variables de compilación. La publishable key es visible para el cliente y no se trata como secreto. La clave de servicio solo se obtiene en la Edge Function desde el entorno de Supabase y no debe incluirse en Flutter, el repositorio ni las capturas. No se reproducen valores de claves en este documento.')

heading('2.8 Pruebas', 1)
paragraph('Se ejecutó una prueba automatizada contra la URL real de salud el 03/10/2026. El script scripts/verify_salud.ps1 compara los métodos GET, HEAD y POST con sus respuestas esperadas y devuelve error cuando un caso falla. La apertura de la web publicada también se observó en navegador. Los casos que requieren cuentas ficticias y políticas RLS activas permanecen pendientes. La matriz completa y el comando reproducible están en docs/PRUEBAS_Y_EVIDENCIAS.md.')

table = doc.add_table(rows=1, cols=5)
table.style = doc.tables[12].style
for cell, value in zip(table.rows[0].cells, ['ID', 'Escenario evaluado', 'Resultado esperado', 'Resultado observado', 'Estado']):
    cell.text = value
cases = [
    ('CP-01', 'Apertura del panel web', 'Página pública disponible', 'Página EIVET observada en Netlify; flujos internos sin probar', 'Aprobado solo para disponibilidad'),
    ('CP-02', 'Salud GET', 'HTTP 200; status online', 'HTTP 200; status online', 'Aprobado'),
    ('CP-03', 'Salud HEAD', 'HTTP 200; sin cuerpo', 'HTTP 200; sin cuerpo', 'Aprobado'),
    ('CP-04', 'Salud POST', 'HTTP 405', 'HTTP 200; status online', 'Fallido'),
    ('CP-05', 'Aislamiento entre dos propietarios', 'Sin lectura ni edición cruzada', 'No ejecutado en producción', 'Pendiente'),
    ('CP-06', 'Permisos de mascotas por rol', 'Personal escribe; propietario solo lee', 'No ejecutado de extremo a extremo', 'Pendiente'),
    ('CP-07', 'Prueba Flutter del repositorio demo', 'Aserciones satisfactorias', 'No se obtuvo salida final de la ejecución', 'Pendiente'),
]
for case in cases:
    for cell, value in zip(table.add_row().cells, case):
        cell.text = value
for row in table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.size = Pt(8)
table.autofit = True
anchor.addprevious(table._element)
paragraph('Tabla 14. Casos de prueba ejecutados y pendientes al 03/10/2026. Fuente: elaboración propia.')
paragraph('Las capturas fechadas de la página y de la respuesta de salud aún deben adjuntarse. Las capturas de permisos y acceso entre cuentas solo se incorporarán después de ejecutar las pruebas con datos ficticios, sin exponer credenciales ni datos clínicos. La respuesta online de salud acredita la consulta de comprobación a PostgreSQL, no el funcionamiento de todos los módulos.')

# Corrige la referencia al proveedor de despliegue, ya decidido y publicado.
for row in doc.tables[12].rows:
    if row.cells[0].text.strip() == 'Despliegue':
        row.cells[1].text = 'Flutter Web en Netlify; APK Android disponible mediante compilación local'
        row.cells[2].text = 'Publicar el panel web por HTTPS y generar la aplicación Android instalable.'
        row.cells[3].text = 'Netlify publica la versión web desde GitHub; responde en https://eivetveterinaria.netlify.app/. La distribución del APK aún requiere evidencia.'

doc.paragraphs[-1].text = ('Pendientes antes de la entrega: medir el proceso actual; adjuntar capturas fechadas del tablero y de producción; verificar el acceso docente al repositorio; registrar la versión exacta de Flutter; confirmar la aplicación de las políticas RLS y ejecutar las pruebas con dos propietarios ficticios. La ciudad del proyecto es Tarija.')
doc.save(OUTPUT)
print(OUTPUT.resolve())
